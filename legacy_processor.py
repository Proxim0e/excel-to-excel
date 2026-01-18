#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
legacy_processor.py

Скрипт для синхронизации "лотов" из родительского Excel-файла в копии шаблона Excel,
скрапинга страниц лотов с сайта и автоматического заполнения данных по участникам.

Ключевые возможности:
- Создаёт листы "lot N" на основе шаблона sample_model.xlsx и записывает в них:
  - заголовок (B1) в формате "Lot nr. N {denumire}",
  - спецификацию (B11) при наличии,
  - прочие поля из TARGET_MAP (D8, C9 и т.д.).
- Первый участник записывается в колонку E, последующие — вставляются как новые колонки
  перед колонкой "Note" с копированием стилей из колонки E (чтобы "Note" оставалась последней).
- Для каждой колонки ставятся формулы:
    <col>7 = =<col>8/A1
    <col>9 = =<col>8/D8
  (формулы записываются непосредственно в Excel; вычисления выполнит Excel при открытии).
- Если участников нет — в E11 добавляется "Achiziţia nu a avut loc" с тёмно-синей заливкой
  и белым шрифтом; вкладка листа помечается тем же цветом.
- Если (price / D8) >= PERCENT_THRESHOLD (по умолчанию 1.3) — ячейка процента подсвечивается
  красным; если ВСЕ участники лота имеют >= порога — вкладка помечается красным.
- Лоты скрапятся параллельно (ThreadPoolExecutor).
- Скрипт определяет максимальный уже существующий номер лота (max_existing). Поведение:
  - APPEND_TO_EXISTING = False (по умолчанию): НЕ изменяем листы с номерами <= max_existing.
    Но для лотов с номерами > max_existing (созданных в этом запуске или отсутствовавших в шаблоне)
    допускается дополнять спецификацию при дублирующихся записях внутри одного запуска.
  - APPEND_TO_EXISTING = True: разрешено дополнять существующие листы (старое поведение).
- Tender URL (страница тендера, откуда берутся ссылки на лоты) берётся:
  - из родительского файла — колонка, содержащая "site"/"url"/"link" (первая непустая ссылка);
  - если не найдена — используется DEFAULT_TENDER_URL.

Автор: адаптация Proxim0e
Дата: 2025-11-20

Примечание о безопасности:
- Скрипт записывает и изменяет файлы в папке RESOURCE_DIR; перед запуском убедитесь в наличии
  резервных копий, особенно шаблона sample_model.xlsx и родительского файла.
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from typing import List, Optional, Set, Tuple, Dict
from urllib.parse import urljoin
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
from copy import copy
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# Конфигурация (настраивайте при необходимости)
# ---------------------------------------------------------------------------
RESOURCE_DIR: str = './resources'                       # папка с входными/выходными файлами
TEMPLATE_FILE: str = os.path.join(RESOURCE_DIR, 'sample_model.xlsx')

# Если в parent-файле нет колонки "site" с URL, используется этот запасной URL
DEFAULT_TENDER_URL: str = "https://achizitii.md/ro/public/tender/21463176"

BASE: str = "https://achizitii.md"
# Шаблон регулярного выражения для распознавания ссылок на лоты
LOT_LINK_RE = re.compile(r'/ro/public/tender/\d+/lot/\d+/?$')

HEADERS: Dict[str, str] = {"User-Agent": "Mozilla/5.0 (compatible; scraper/1.0; +https://example.com/bot)"}

MAX_WORKERS: int = 16          # параллелизм для скачивания страниц лотов

# Поведение при существующих листах:
# False (по умолчанию) — НЕ добавляем новые данные в листы с lot <= max_existing,
# True  — дополняем существующие листы (старое поведение)
APPEND_TO_EXISTING: bool = False

# Список возможных заголовков колонок в родительском файле (можно расширить)
COLS = {
    'nr_lot':            ['Nr. Lot', 'Nr Lot', 'Nr lot', 'Nr. lot'],
    'denumire':          ['Denumirea Lotului\n04.09.2025', 'Denumirea Lotului', 'Denumire Lot NEW2'],
    'specificatie':      ['Specificația Tehnică\n04.09.2025', 'Specificația Тehnică', 'Specificarea техническая NEW2'],
    'unitate_masura':    ['Unitatea de măsură'],
    'cantitate_total':   ['Cantitatea Totală'],
    'suma_alocata':      ['Suma alocată'],
}

# Куда писать определённые поля в шаблоне
TARGET_MAP = {
    "denumire": ("B1",),
    "specificatie": ("B11",),
    "unitate_masura": ("C9",),
    "cantitate_total": ("A1",),
    "suma_alocata": ("D8",)
}

# Оформление: цвета и шрифты
NOT_HELD_FILL_DARK = PatternFill(start_color="001f4d", end_color="001f4d", fill_type="solid")  # тёмно-синий
NOT_HELD_FONT_WHITE = Font(color="FFFFFF", bold=True)

HIGH_PERCENT_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")  # красный
HIGH_PERCENT_FONT = Font(color="000000", bold=True)

PERCENT_THRESHOLD: float = 1.3  # 130%


# ---------------------------------------------------------------------------
# Вспомогательные функции
# ---------------------------------------------------------------------------
def find_file(prefix: str, ext: str = '.xlsx') -> str:
    """
    Найти файл в RESOURCE_DIR по началу имени (prefix).
    Возвращает полный путь к первому совпавшему файлу.
    Выбрасывает FileNotFoundError, если файл не найден.
    """
    for fname in os.listdir(RESOURCE_DIR):
        if fname.startswith(prefix) and fname.endswith(ext):
            return os.path.join(RESOURCE_DIR, fname)
    raise FileNotFoundError(f"Файл с префиксом '{prefix}' не найден в директории {RESOURCE_DIR}!")


def safe_col_idx(names: List[str], headers: List[Optional[str]]) -> Optional[int]:
    """
    Найти индекс (1-based) колонки в headers по списку возможных имён names.
    Сравнение регистронезависимое; возвращает None, если не найдено.
    """
    if isinstance(names, str):
        names = [names]
    for name in names:
        if name is None:
            continue
        for i, h in enumerate(headers):
            if h is None:
                continue
            try:
                if str(h).strip().lower() == str(name).strip().lower():
                    return i + 1
            except Exception:
                continue
    return None


def safe_row_get(row: Tuple, keys: List[str], headers: List[Optional[str]]):
    """
    Получить значение из строки row по одной из возможных названий в keys.
    Возвращает значение или None.
    """
    idx = safe_col_idx(keys, headers)
    if idx is not None:
        return row[idx - 1]
    return None


def get_soup(url: str) -> Optional[BeautifulSoup]:
    """
    Скачать страницу по URL и распарсить через BeautifulSoup с парсером html.parser.
    Возвращает объект BeautifulSoup или None при ошибке.
    """
    try:
        resp = requests.get(url, headers=HEADERS, timeout=20)
        resp.raise_for_status()
        return BeautifulSoup(resp.text, "html.parser")
    except Exception as e:
        print(f"[ERROR] Ошибка загрузки {url}: {e}")
        return None


def extract_lot_links(tender_url: str) -> List[str]:
    """
    Собрать список ссылок на лоты со страницы тендера.
    Возвращает отсортированный список уникальных абсолютных URL.
    """
    soup = get_soup(tender_url)
    if not soup:
        return []
    links: Set[str] = set()
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        full = urljoin(BASE, href)
        if LOT_LINK_RE.search(href) or LOT_LINK_RE.search(full.replace(BASE, "")):
            links.add(full.split('#')[0].rstrip("/"))
    return sorted(links)


def strip_price_label(s: str) -> str:
    """
    Очищает префиксы вроде 'prețul ofertei: ' из текстовых представлений цены.
    """
    if not s:
        return s
    s = s.strip()
    s = re.sub(r'(?i)prețul ofertei[:\s]*', '', s)
    s = re.sub(r'(?i)preţul ofertei[:\s]*', '', s)
    s = re.sub(r'(?i)preţul[:\s]*', '', s)
    s = re.sub(r'(?i)preț[:\s]*', '', s)
    s = re.sub(r'(?i)preţ[:\s]*', '', s)
    s = re.sub(r'(?i)pre[uț]ț?[:\s]*', '', s)
    return s.strip()


def parse_lot_html(soup: BeautifulSoup) -> Tuple[Optional[str], List[Tuple[str, str]]]:
    """
    Разбор HTML страницы лота:
    - ищем заголовок лота;
    - собираем список участников: (название, цена_как_строка).
    Возвращает (title, participants).
    """
    if not soup:
        return None, []

    title: Optional[str] = None
    for sel in ("h1", ".tender__page__title h1", ".tender__page__title"):
        el = soup.select_one(sel)
        if el and el.get_text(strip=True):
            title = el.get_text(strip=True)
            break
    if not title:
        mt = soup.find("title")
        title = mt.get_text(strip=True) if mt else "(untitled)"

    participants: List[Tuple[str, str]] = []
    infos = soup.select(".participant-container-body-info")
    if infos:
        for info in infos:
            name = ""
            nnode = info.select_one(".participant-container-body-item.participant-title, .participant-title, .participant-title-div")
            if nnode:
                name = nnode.get_text(separator=" ", strip=True)
            price = ""
            pnode = info.select_one(".participant-container-body .participant-price, .participant-price")
            if pnode:
                price_text = pnode.get_text(separator=" ", strip=True)
                price = strip_price_label(price_text)
            if name or price:
                participants.append((name, price))
    else:
        # fallback: собрать отдельно title_nodes и price_nodes, если структура иная
        title_nodes = soup.select(".participant-container-body .participant-title, .participant-title")
        price_nodes = soup.select(".participant-container-body .participant-price, .participant-price")
        if title_nodes and price_nodes and len(title_nodes) == len(price_nodes):
            for t, p in zip(title_nodes, price_nodes):
                price_clean = strip_price_label(p.get_text(strip=True))
                participants.append((t.get_text(strip=True), price_clean))

    return title, participants


def parse_lot_page(lot_url: str) -> Dict:
    """
    Утилита для использования в пуле воркеров: загружает lot_url, парсит и возвращает словарь с данными.
    Формат возвращаемого словаря:
      {'url': lot_url, 'title': title_or_None, 'participants': [(name, price_str), ...], 'error': None_or_str}
    """
    try:
        soup = get_soup(lot_url)
        if not soup:
            return {'url': lot_url, 'title': None, 'participants': [], 'error': 'load_failed'}
        title, participants = parse_lot_html(soup)
        return {'url': lot_url, 'title': title, 'participants': participants, 'error': None}
    except Exception as e:
        return {'url': lot_url, 'title': None, 'participants': [], 'error': str(e)}


def clean_company_name(raw_name: Optional[str]) -> Optional[str]:
    """
    Простая нормализация названия участника.
    Убирает префиксы вроде 'Denumirea participantului:' и обрывает строку перед 'Preţ' и т.п.
    """
    if not raw_name:
        return raw_name
    s = raw_name.strip()
    s = re.sub(r'(?i)denumirea participantului[:\s]*', '', s).strip()
    s = re.sub(r'^\s*Denumirea[:\s]*', '', s, flags=re.I).strip()
    s = re.split(r'\s+Preţ|Preţul|Preț', s)[0].strip()
    return s


def parse_price_to_number(price_str: Optional[str]):
    """
    Попытка извлечь число из строки цены. Поддерживаются форматы:
    - "30 833,35 MDL" -> 30833.35
    - "30.833,35" -> 30833.35
    - "30833.35" -> 30833.35
    - "3 083 335" -> 3083335
    Возвращает int/float или None.
    """
    if not price_str:
        return None
    s = price_str.strip()
    m = re.search(r'[\d\s\.,]+', s)
    if not m:
        return None
    s = m.group(0).strip()
    s = s.replace('\u00A0', ' ')
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '')
            s = s.replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        s = s.replace(' ', '')
        s = s.replace(',', '.')
    else:
        s = s.replace(' ', '')
    s = re.sub(r'[^\d\.]', '', s)
    if not s:
        return None
    try:
        if '.' in s:
            return float(s)
        else:
            return int(s)
    except Exception:
        return None


def clean_denumire_for_B1(denumire: Optional[str]) -> str:
    """
    Очищает текст 'denumire' перед помещением в B1:
    убирает префиксы 'Lot nr N...' и лишние пробелы.
    """
    if not denumire:
        return ""
    s = str(denumire).strip()
    s = re.sub(r'(?i)^\s*Lot(?:ul)?\s*nr\.?\s*\d+\s*[-–:\)]*\s*', '', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s


def normalize_spaces(s: Optional[str]) -> str:
    """
    Приводит последовательности пробельных символов к одному пробелу.
    """
    if not s:
        return ""
    return re.sub(r'\s+', ' ', str(s)).strip()


def copy_column_styles(ws: Worksheet, src_col_idx: int, dst_col_idx: int, max_row: int) -> None:
    """
    Копирует стиль, формат и ширину из src_col_idx в dst_col_idx для строк от 1 до max_row.
    Не копирует значения.
    """
    src_letter = get_column_letter(src_col_idx)
    dst_letter = get_column_letter(dst_col_idx)

    # Ширина колонки
    try:
        src_dim = ws.column_dimensions.get(src_letter)
        if src_dim and src_dim.width is not None:
            ws.column_dimensions[dst_letter].width = src_dim.width
    except Exception:
        pass

    for r in range(1, max_row + 1):
        src_cell = ws.cell(row=r, column=src_col_idx)
        dst_cell = ws.cell(row=r, column=dst_col_idx)

        # Стиль и формат (без устаревшего .copy())
        try:
            dst_cell.font = copy(src_cell.font)
        except AttributeError:
            pass
        try:
            dst_cell.fill = copy(src_cell.fill)
        except AttributeError:
            pass
        try:
            dst_cell.border = copy(src_cell.border)
        except AttributeError:
            pass
        try:
            dst_cell.number_format = src_cell.number_format
        except AttributeError:
            pass
        try:
            dst_cell.alignment = copy(src_cell.alignment)
        except AttributeError:
            pass
        try:
            dst_cell.protection = copy(src_cell.protection)
        except AttributeError:
            pass


# ---------------------------------------------------------------------------
# Основной рабочий процесс
# ---------------------------------------------------------------------------
def main() -> None:
    """
    Главная функция:
    - читает родительский файл (ищет файл в RESOURCE_DIR по префиксу '-');
    - определяет существующие листы lot N в шаблоне;
    - создаёт новые листы и/или дополняет (в зависимости от APPEND_TO_EXISTING и max_existing);
    - собирает ссылки на лоты из тендера (TENDER URL берётся либо из parent-файла, либо DEFAULT_TENDER_URL);
    - параллельно парсит страницы лотов и записывает участников в соответствующие листы;
    - сохраняет итоговую книгу в RESOURCE_DIR.
    """
    try:
        # Найти родительский файл в RESOURCE_DIR (ищем файл, начинающийся с '-')
        parent_file = find_file('-')
        print(f"Родительский файл: {parent_file}")
        print(f"Шаблонный файл: {TEMPLATE_FILE}")

        # Загружаем книги
        wb_parent = load_workbook(parent_file, data_only=True)
        ws_parent = wb_parent.active
        wb_child = load_workbook(TEMPLATE_FILE)
        sheet_names = wb_child.sheetnames
        template_sheet = wb_child[sheet_names[0]]

        # Создаём маппинг существующих листов в шаблоне (lot N -> worksheet)
        lot_regex = re.compile(r'^lot (\d+)$')
        parent_lot_sheets: Dict[int, object] = {}
        for sname in wb_child.sheetnames:
            m = lot_regex.match(sname)
            if m:
                try:
                    num = int(m.group(1))
                    parent_lot_sheets[num] = wb_child[sname]
                except AttributeError:
                    continue
        print(f"Найдено в шаблоне {len(parent_lot_sheets)} готовых листов.")

        # Вычисляем максимальный существующий номер лота в шаблоне
        max_existing: int = max(parent_lot_sheets.keys()) if parent_lot_sheets else 0
        print(f"[INFO] Максимальный существующий номер лота в шаблоне: {max_existing}")
        if not APPEND_TO_EXISTING:
            print(f"[INFO] APPEND_TO_EXISTING=False: НЕ изменяю листы с номерами <= {max_existing}. "
                  f"Создаю/обрабатываю только лоты с номера {max_existing + 1} и выше.")

        # Множество номеров лотов, созданных в ходе текущего запуска (позволяет дополнять их, если встречаются дубли)
        created_in_run: Set[int] = set()

        # Заголовки родительского файла
        headers: List[Optional[str]] = [cell.value for cell in ws_parent[1]]
        print("Заголовки в родительском файле:", headers)

        # Ищем колонку с URL тендера ("site" / "url" / "link") и берём первую непустую ссылку
        def find_site_col(headers_list: List[Optional[str]]) -> Optional[int]:
            for i, h in enumerate(headers_list):
                if not h:
                    continue
                hstr = str(h).strip().lower()
                if hstr in ("site", "url", "link"):
                    return i + 1
            for i, h in enumerate(headers_list):
                if not h:
                    continue
                hstr = str(h).strip().lower()
                if "site" in hstr or "url" in hstr or "link" in hstr:
                    return i + 1
            return None

        site_col_idx = find_site_col(headers)
        tender_url = DEFAULT_TENDER_URL
        if site_col_idx:
            for row in ws_parent.iter_rows(min_row=2, values_only=True):
                val = row[site_col_idx - 1]
                if val and isinstance(val, str):
                    v = val.strip()
                    if v.lower().startswith("http"):
                        tender_url = v
                        break
        print(f"[INFO] Использую Tender URL: {tender_url}")

        # Определим индекс колонки "Note" в шаблоне (если есть)
        def find_note_col(ws) -> int:
            for c, cell in enumerate(ws[1], start=1):
                if cell.value and isinstance(cell.value, str) and cell.value.strip().lower() == "note":
                    return c
            return ws.max_column + 1

        note_col_template = find_note_col(template_sheet)
        print(f"[INFO] Note column in template: {get_column_letter(note_col_template)} (index {note_col_template})")

        # ---------------------------
        # Создание / обновление листов на основе родительского файла
        # ---------------------------
        for row in ws_parent.iter_rows(min_row=2, values_only=True):
            nr_lot = safe_row_get(row, COLS['nr_lot'], headers)
            if nr_lot is None or (isinstance(nr_lot, str) and not nr_lot.strip()):
                continue
            try:
                nr_lot_int = int(float(nr_lot))
            except Exception:
                print(f"[WARN] Некорректный номер лота: {nr_lot}")
                continue

            denumire = safe_row_get(row, COLS['denumire'], headers)
            specificatie_data = safe_row_get(row, COLS['specificatie'], headers)

            # Если лист уже есть в parent_lot_sheets
            if nr_lot_int in parent_lot_sheets:
                # Случай: лист присутствовал в шаблоне до запуска (<= max_existing)
                if nr_lot_int <= max_existing:
                    if APPEND_TO_EXISTING:
                        # старое поведение: можно дополнять спецификацию в B11
                        ws_target = parent_lot_sheets[nr_lot_int]
                        if specificatie_data:
                            existing = ws_target["B11"].value or ""
                            add = str(specificatie_data).strip()
                            if add and add not in existing:
                                new_val = (existing + "\n" + add).strip() if existing else add
                                ws_target["B11"] = new_val
                                print(f"[Excel] Дополнена спецификация для lot {nr_lot_int}")
                        # продолжаем на следующую строку
                        continue
                    else:
                        # пропускаем старые листы
                        print(f"[SKIP] Лист для lot {nr_lot_int} уже существует в шаблоне (<= {max_existing}) — пропускаю.")
                        continue
                else:
                    # nr_lot_int > max_existing: разрешаем дополнять (это "новый" лист для нас)
                    ws_target = parent_lot_sheets[nr_lot_int]
                    if specificatie_data:
                        existing = ws_target["B11"].value or ""
                        add = str(specificatie_data).strip()
                        if add and add not in existing:
                            new_val = (existing + "\n" + add).strip() if existing else add
                            ws_target["B11"] = new_val
                            print(f"[Excel] Дополнена спецификация для (нового) lot {nr_lot_int}")
                    continue

            # Иначе — создаём новый лист на основе шаблона и заполняем его
            new_sheet = wb_child.copy_worksheet(template_sheet)
            new_title = f"lot {nr_lot_int}"
            new_sheet.title = new_title
            parent_lot_sheets[nr_lot_int] = new_sheet
            created_in_run.add(nr_lot_int)
            print(f"[Excel] Создаем лист: {new_sheet.title}")

            # Устанавливаем заголовок B1
            den_text_clean = clean_denumire_for_B1(denumire)
            value_to_write = normalize_spaces(f"Lot nr. {nr_lot_int} {den_text_clean}".strip())
            for cell in TARGET_MAP['denumire']:
                new_sheet[cell] = value_to_write
                try:
                    new_sheet[cell].font = Font(bold=True)
                except Exception:
                    pass
            print(f"[EXCEL] Установлен заголовок для листа '{new_sheet.title}' в {TARGET_MAP['denumire']}: {value_to_write!r}")

            # Пишем остальные поля (кроме specificatie)
            for key, target_cells in TARGET_MAP.items():
                if key in ("denumire", "specificatie"):
                    continue
                value = safe_row_get(row, COLS[key], headers)
                if value is not None:
                    for cell in target_cells:
                        new_sheet[cell] = value

            # initial specificatie
            if specificatie_data:
                new_sheet["B11"] = str(specificatie_data).strip()

        # ---------------------------
        # Сбор ссылок на лоты и параллельный парсинг
        # ---------------------------
        print("[SCRAPE] Собираем ссылки на лоты...")
        lot_links = extract_lot_links(tender_url)
        print(f"[SCRAPE] Найдено {len(lot_links)} ссылок на лоты. Параллельность: {MAX_WORKERS} воркеров.")

        results: List[Dict] = []
        if lot_links:
            workers = min(MAX_WORKERS, max(2, len(lot_links)))
            with ThreadPoolExecutor(max_workers=workers) as ex:
                future_to_url = {ex.submit(parse_lot_page, url): url for url in lot_links}
                for fut in as_completed(future_to_url):
                    url = future_to_url[fut]
                    try:
                        res = fut.result()
                    except Exception as e:
                        print(f"[ERROR] Ошибка при обработке {url}: {e}")
                        res = {'url': url, 'title': None, 'participants': [], 'error': str(e)}
                    results.append(res)
                    if res.get('error'):
                        print(f"[SCRAPE] {url} -> error: {res['error']}")
                    else:
                        tcount = len(res.get('participants') or [])
                        title_sample = normalize_spaces(res.get('title') or "")
                        print(f"[SCRAPE] {url} -> title: {title_sample!r}, participants: {tcount}")

        # ---------------------------
        # Применяем результаты к Excel (запись участников)
        # ---------------------------
        for r in results:
            lot_url = r.get('url')
            title = r.get('title')
            participants = r.get('participants') or []
            lot_number: Optional[int] = None
            if title:
                m = re.search(r'Lot(?:ul)?\s*nr\.?\s*\.*\s*(\d+)', title, flags=re.I)
                if m:
                    lot_number = int(m.group(1))
            if lot_number is None:
                print(f"[WARN] Не удалось извлечь номер лота из title для URL {lot_url!r} title={title!r}. Пропускаю.")
                continue

            # Пропускаем обновление участников для старых лотов (если APPEND_TO_EXISTING == False)
            if (not APPEND_TO_EXISTING) and lot_number <= max_existing:
                print(f"[SKIP] Пропускаю обновление участников для lot {lot_number} — этот лист существовал в шаблоне (<= {max_existing}).")
                continue

            if lot_number not in parent_lot_sheets:
                print(f"[WARN] Для lot {lot_number} нет листа в книге (пропуск обновления участников).")
                continue

            ws = parent_lot_sheets[lot_number]

            # Поиск колонки Note для данного листа
            def find_note_col_ws(ws_loc) -> int:
                for c, cell in enumerate(ws_loc[1], start=1):
                    if cell.value and isinstance(cell.value, str) and cell.value.strip().lower() == "note":
                        return c
                return ws_loc.max_column + 1

            note_col = find_note_col_ws(ws)
            if note_col <= 5:
                note_col = 6

            # Если участников нет — пометим E11 и продолжим
            if not participants:
                print(f"[EXCEL] lot {lot_number}: участников не найдено — вставляем 'Achiziţia nu a avut loc' в E11.")
                ws["E11"] = "Achiziţia nu a avut loc"
                ws["E11"].fill = NOT_HELD_FILL_DARK
                ws["E11"].alignment = Alignment(horizontal="center", vertical="center")
                try:
                    ws["E11"].font = NOT_HELD_FONT_WHITE
                except Exception:
                    pass
                try:
                    ws.sheet_properties.tabColor = "001f4d"
                except Exception:
                    pass
                continue

            # Записываем участников: первый в E, остальные — новые колонки перед Note
            max_row = ws.max_row if ws.max_row > 1 else 50
            src_col = 5  # E
            high_flags: List[bool] = []

            for idx, (raw_name, raw_price) in enumerate(participants, start=1):
                if idx == 1:
                    target_col = src_col
                else:
                    insert_at = note_col
                    ws.insert_cols(insert_at, amount=1)
                    target_col = insert_at
                    copy_column_styles(ws, src_col, target_col, max_row)
                    note_col += 1

                col_letter = get_column_letter(target_col)
                name = clean_company_name(raw_name) or "(без имени)"
                price_value = parse_price_to_number(raw_price) if raw_price else None

                # Ячейки для записи
                cell_name = f"{col_letter}2"
                cell_price = f"{col_letter}8"
                cell_formula_divA1 = f"{col_letter}7"
                cell_formula_divD8 = f"{col_letter}9"

                ws[cell_name] = f"Ofertant: {name}"
                try:
                    ws[cell_name].font = Font(bold=True)
                except Exception:
                    pass

                # Записываем цену как число (если распознали)
                if price_value is not None:
                    ws[cell_price] = price_value
                else:
                    ws[cell_price] = ""

                # Ставим формулы
                ws[cell_formula_divA1] = f"={col_letter}8/A1"
                ws[cell_formula_divD8] = f"={col_letter}8/D8"

                # Вычислим процент и применим выделение, если нужно
                is_high = False
                try:
                    d8_raw = ws["D8"].value
                    d8_val = None
                    if isinstance(d8_raw, (int, float)):
                        d8_val = float(d8_raw)
                    else:
                        d8_val = parse_price_to_number(str(d8_raw)) if d8_raw is not None else None
                    if price_value is not None and d8_val not in (None, 0):
                        percent = float(price_value) / float(d8_val)
                        if percent >= PERCENT_THRESHOLD:
                            is_high = True
                            ws[cell_formula_divD8].fill = HIGH_PERCENT_FILL
                            ws[cell_formula_divD8].font = HIGH_PERCENT_FONT
                            ws[cell_formula_divD8].alignment = Alignment(horizontal="center", vertical="center")
                except Exception:
                    pass

                high_flags.append(is_high)
                print(f"[EXCEL] lot {lot_number}: записан участник в {col_letter} -> '{name}' / price={price_value} high={is_high}")

            # Если ВСЕ участники high — пометить вкладку красным
            try:
                if high_flags and all(high_flags):
                    ws.sheet_properties.tabColor = "FF0000"
                    print(f"[EXCEL] lot {lot_number}: все участники >={int(PERCENT_THRESHOLD*100)}% -> вкладка помечена красным")
            except Exception:
                pass

        # ---------------------------
        # Сохранение итоговой книги
        # ---------------------------
        parent_basename = os.path.basename(parent_file).replace('-', '')
        timestamp = datetime.now().strftime("%d.%m.%Y_%H-%M-%S")
        output_name = f"{RESOURCE_DIR}/{os.path.splitext(parent_basename)[0]}_{timestamp}_upd.xlsx"
        wb_child.save(output_name)
        print(f"\nГотово! Файл сохранён как {output_name}")

    except FileNotFoundError as e:
        print(e)
    except Exception as e:
        print(f"Неизвестная ошибка: {e}")


if __name__ == "__main__":
    main()