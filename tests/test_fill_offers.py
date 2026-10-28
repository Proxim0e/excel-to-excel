import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.anexa22 import evidence as ev
from src.excel import fill_offers as fo
from src.manifest import MANIFEST_NAME, migrate

TEMPLATE_ROW3 = "Producator: \nȚara: \nModel:"


class EvidenceTests(unittest.TestCase):
    def test_model_tokens(self):
        self.assertEqual(ev.model_tokens("Electrod lama (Codul produsului:F4141)"), ["F4141"])
        self.assertEqual(ev.model_tokens("M3D, 3011-00010-01"), ["3011-00010-01"])
        self.assertEqual(ev.model_tokens("Masca adult"), [])

    def test_spec_values(self):
        values = ev.spec_values("lungime 70 mm (+-5mm), Ø 49mm, 0~300mm Hg, 500 ml, 3m, 1,5 ml")
        self.assertEqual(values[:6], [("70", "mm"), ("5", "mm"), ("49", "mm"), ("300", "mm"), ("500", "ml"),
                                      ("3", "m")])
        self.assertIn(("1.5", "ml"), values)

    def test_format(self):
        items = [ev.Evidence("catalog.pdf", 3, ['49 mm diameter', 'No. 5270'])]
        self.assertEqual(ev.format_evidence(items), '// "49 mm diameter" "No. 5270", p[3], catalog.pdf')


class FillTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        conn = sqlite3.connect(self.root / MANIFEST_NAME)
        migrate(conn)
        conn.execute("INSERT INTO documents (sha256, size, first_seen, doc_type) VALUES "
                     "('a22', 1, 'now', 'anexa22'), ('cat', 1, 'now', 'datasheet')")
        for sha, name in (("a22", "anexa22.pdf"), ("cat", "catalog.pdf")):
            conn.execute("INSERT INTO document_refs (path, tender_id, lot_no, lot_title, participant, original_name,"
                         " url, sha256, status, recorded_at) VALUES (?, 'T', 2, 'Lot 2', 'SOGNO', ?, 'u', ?, 'ok', 'now')",
                         (f"lot2/{name}", name, sha))
        conn.execute("INSERT INTO document_pages (sha256, page_no, text, page_type) VALUES "
                     "('a22', 1, 'Anexa 22 Metpak 5270 500 ml', 'anexa22'), "
                     "('cat', 1, 'Riester Metpak\nNo. 5270\nfor bags from 500 ml\nmanometer 49 mm diameter', 'datasheet')")
        conn.execute("INSERT INTO anexa22_offers (sha256, page_no, participant, lot_no, model, country, producer, "
                     "spec_offered, standards, dm, parsed_at) VALUES ('a22', 1, 'SOGNO', 2, 'Metpak 500ml (cod 5270)', "
                     "'Germania', 'RIESTER', 'Manseta 500 ml, Manometru: Ø 49mm', 'CE, ISO', '', 'now')")
        conn.commit()
        conn.close()

        wb = Workbook()
        ws = wb.active
        ws.title = "lot 2"
        ws["E2"], ws["E3"] = "Ofertant: SOGNO", TEMPLATE_ROW3
        ws["F2"], ws["F4"] = "Ofertant: Labromed Laborator SRL", "уже заполнено человеком"
        ws["G2"] = "Ofertant: "
        self.xlsx = self.root / "result.xlsx"
        wb.save(self.xlsx)

    def tearDown(self):
        self.tmp.cleanup()

    def test_fill_rows_and_evidence(self):
        stats = fo.fill_result(self.xlsx, self.root)
        self.assertEqual((stats["participants"], stats["matched"], stats["missing"]), (2, 1, 1))
        ws = load_workbook(self.xlsx)["lot 2"]
        self.assertEqual(ws["E3"].value, "Producator: RIESTER\nȚara: Germania\nModel: Metpak 500ml (cod 5270)")
        self.assertEqual(ws["E4"].value, "CE, ISO")
        spec = ws["E11"].value
        self.assertTrue(spec.startswith("Manseta 500 ml, Manometru: Ø 49mm\n// "))
        self.assertIn('"No. 5270"', spec)
        self.assertIn('"manometer 49 mm diameter"', spec)
        self.assertIn("p[1], catalog.pdf", spec)
        self.assertNotIn("anexa22.pdf", spec)  # Anexa 22 не подтверждает сама себя
        self.assertEqual(ws["F4"].value, "уже заполнено человеком")

    def test_nothing_overwritten_on_second_run(self):
        fo.fill_result(self.xlsx, self.root)
        wb = load_workbook(self.xlsx)
        wb["lot 2"]["E11"] = "правка человека"
        wb.save(self.xlsx)
        stats = fo.fill_result(self.xlsx, self.root)
        self.assertEqual((stats["row3"], stats["row11"]), (0, 0))
        self.assertEqual(load_workbook(self.xlsx)["lot 2"]["E11"].value, "правка человека")

    def test_dm_cell_prefers_dm_numbers(self):
        rec = {"dm": ["DM000466393,DM000466394"], "standards": ["CE/DM000466393"]}
        self.assertEqual(fo.dm_cell(rec), "DM000466393\nDM000466394")


if __name__ == "__main__":
    unittest.main()
