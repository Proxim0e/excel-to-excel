"""
Тесты проверки AMDM. Примеры строк взяты из реальных книг оценки (CDM27_B, CPAS27-P1).
Запуск: python -m unittest discover -s tests -t .   (или python -m pytest)
"""
import io
import tempfile
import unittest
import zipfile
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

from src.amdm.cache import RegistryCache
from src.amdm.fetch import CsvImporter, header_field
from src.amdm.matcher import check_offer, model_matches, overall_status, producer_matches
from src.amdm.models import OK, CHECK, FAIL, NO_DM, IN_PROGRESS, OfferRegistration, RegistryRecord
from src.amdm.parse import normalize_dm, parse_row3, parse_row4
from src.amdm.report import SHEET, write_report
from src.amdm.workbook import load_workbook_tolerant, read_offers


class ParseRow4Test(unittest.TestCase):

    def test_dm_order_date(self):
        p = parse_row4("DM000365870 Rg04-000183 05-08-2022/ pentru DM")
        self.assertEqual(p["dm_numbers"], ["DM000365870"])
        self.assertEqual(p["order"], "Rg04-000183")
        self.assertEqual(p["order_date"], "2022-08-05")
        self.assertIn("pentru DM", p["note"])

    def test_line_break_and_several_numbers(self):
        p = parse_row4("DM000377731\nRg04-000243 17-10-2022")
        self.assertEqual((p["dm_numbers"], p["order_date"]), (["DM000377731"], "2022-10-17"))
        self.assertEqual(parse_row4("DM000123456 DM000654321")["dm_numbers"], ["DM000123456", "DM000654321"])

    def test_quality_documents_only(self):
        self.assertEqual(parse_row4("CE, ISO 13485")["quality_docs"], ["CE", "ISO 13485"])
        p = parse_row4("CE/ISO ")
        self.assertEqual((p["dm_numbers"], p["quality_docs"]), ([], ["CE", "ISO"]))

    def test_empty(self):
        for v in (None, "None", ""):
            self.assertEqual(parse_row4(v)["dm_numbers"], [])

    def test_normalize_dm(self):
        self.assertEqual(normalize_dm("dm 655298"), "DM000655298")
        self.assertEqual(normalize_dm("CE/DM000232508"), "DM000232508")
        self.assertIsNone(normalize_dm("ISO 9001"))


class ParseRow3Test(unittest.TestCase):

    def test_wrapped_values_are_joined(self):
        p = parse_row3("Producator: BSS MEDICAL\nSUPPLY CO.,\nLIMITED\nȚara: China\n"
                       "Model:  Electrod\nECG 8 ore\n(Codul\nprodusului:B3644)")
        self.assertEqual(p["producer"], "BSS MEDICAL SUPPLY CO., LIMITED")
        self.assertEqual(p["country"], "China")
        self.assertEqual(p["model"], "Electrod ECG 8 ore (Codul produsului:B3644)")

    def test_template_and_empty(self):
        self.assertEqual(parse_row3("Producator: \nȚara: \nModel:"), {"producer": None, "country": None, "model": None})
        self.assertEqual(parse_row3("None")["producer"], None)

    def test_old_template_labels(self):
        p = parse_row3("Producator:  Natus\nȚara de Origine: Irlanda \nModel art: 25V26")
        self.assertEqual((p["producer"], p["country"], p["model"]), ("Natus", "Irlanda", "25V26"))


class MatchingTest(unittest.TestCase):

    def test_producer(self):
        self.assertTrue(producer_matches("Mindray", "Shenzhen Mindray Bio-Medical Electronics Co., Ltd")[0])
        self.assertTrue(producer_matches("B.Braun", "B. Braun Melsungen AG")[0])
        self.assertTrue(producer_matches("Intco Medical", "Shandong Intco Medical Products Co., Ltd.")[0])
        self.assertTrue(producer_matches("P.J. Dahlhausen & Co. GmbH", "P.J. DAHLHAUSEN & CO. GMBH")[0])
        self.assertFalse(producer_matches("Fiab", "Intersurgical Ltd")[0])
        self.assertFalse(producer_matches(None, "Intersurgical")[0])

    def test_model_both_directions(self):
        rec = RegistryRecord(dm="DM000078224", catalog="8250731SP", model="Infusomat Space Line")
        self.assertTrue(model_matches("Infusomat Space Line Neutrapur Ref.N8250731SP", rec)[0])
        rec2 = RegistryRecord(dm="DM1", catalog="630-240-00B, 630-240-05B")
        self.assertTrue(model_matches("630 240 05B", rec2)[0])
        self.assertFalse(model_matches("T20163", rec2)[0])


def offer(**kw):
    base = dict(sheet="lot 1", column=5, supplier="Labromed", producer="P.J. Dahlhausen & Co. GmbH",
                model="43.018.00.013", dm_numbers=["DM000654009"], order="Rg04-000259", order_date="2023-11-09")
    base.update(kw)
    return OfferRegistration(**base)


RECORD = RegistryRecord(dm="DM000654009", name="Cateter", producer="P.J. Dahlhausen & Co. GmbH",
                        catalog="43.018.00.013, 43.018.00.016", order="Rg04-000259", order_date="2023-11-09")


class CheckOfferTest(unittest.TestCase):

    def test_green(self):
        [r] = check_offer(offer(), {"DM000654009": RECORD}, {"DM000654009"})
        self.assertEqual(r.status, OK)

    def test_not_found_is_red(self):
        [r] = check_offer(offer(), {"DM000654009": None}, {"DM000654009"})
        self.assertEqual(r.status, FAIL)
        [r] = check_offer(offer(), {"DM000654009": None}, set())
        self.assertIn("перепроверить", r.reasons[0])

    def test_mismatch_is_yellow(self):
        [r] = check_offer(offer(producer="Fiab"), {"DM000654009": RECORD}, set())
        self.assertEqual(r.status, CHECK)
        self.assertIn("Производитель", r.reasons[0])
        [r] = check_offer(offer(model="X-99999"), {"DM000654009": RECORD}, set())
        self.assertEqual(r.status, CHECK)
        [r] = check_offer(offer(order_date="2023-11-10"), {"DM000654009": RECORD}, set())
        self.assertEqual(r.status, CHECK)

    def test_not_fetched_and_no_dm(self):
        [r] = check_offer(offer(), {}, set())
        self.assertEqual(r.status, IN_PROGRESS)
        [r] = check_offer(offer(dm_numbers=[], quality_docs=["ISO 13485"]), {}, set())
        self.assertEqual(r.status, CHECK)
        [r] = check_offer(offer(dm_numbers=[]), {}, set())
        self.assertEqual(r.status, NO_DM)
        self.assertEqual(check_offer(offer(dm_numbers=[], producer=None, model=None), {}, set()), [])

    def test_kit_overall_is_worst(self):
        rs = check_offer(offer(dm_numbers=["DM000654009", "DM000000001"]),
                         {"DM000654009": RECORD, "DM000000001": None}, {"DM000000001"})
        self.assertEqual(overall_status(rs), FAIL)


class CacheTest(unittest.TestCase):

    def test_put_get_and_not_found_ttl(self):
        with tempfile.TemporaryDirectory() as d:
            c = RegistryCache(Path(d) / "r.sqlite")
            c.put(RECORD, source="test")
            self.assertEqual(c.get("DM000654009").catalog, RECORD.catalog)
            self.assertEqual(c.latest_order_date(), "2023-11-09")
            c.mark_not_found("DM000000001")
            self.assertTrue(c.not_found_fresh("DM000000001"))
            self.assertFalse(c.not_found_fresh("DM000000001", now=datetime.now() + timedelta(days=8)))
            c.put(RegistryRecord(dm="DM000000001"), source="test")       # нашёлся позже
            self.assertFalse(c.not_found_fresh("DM000000001"))
            c.close()


class CsvImportTest(unittest.TestCase):

    def test_semicolon_csv_with_site_headers(self):
        text = ("Nr;Denumire;Den.comerc.;Model;Nr. catalog;Tara;Producatorul;Reprezentant;Ordin;Data;Cod vamal\n"
                "DM000927883;Cateter;Cat-X;M1;123-45;Germania;B. Braun;Rep SRL;Rg04-305;06-10-2026;9018\n")
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "r.csv"
            p.write_text(text, encoding="utf-8")
            [rec] = CsvImporter().read(p)
        self.assertEqual((rec.dm, rec.trade_name, rec.catalog, rec.order_date),
                         ("DM000927883", "Cat-X", "123-45", "2026-10-06"))

    def test_header_names(self):
        self.assertEqual(header_field("Producătorul"), "producer")
        self.assertEqual(header_field("Den.comerc."), "trade_name")
        self.assertEqual(header_field("Nr. catalog"), "catalog")
        self.assertEqual(header_field("Nr"), "dm")


def lot_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "lot 23"
    ws["E2"] = "Ofertant: \"ALMICOM-TECH\" S.R.L."
    ws["E3"] = "Producator: Emed\nȚara: Polonia\nModel: 630-240-00B "
    ws["E4"] = "CE/ISO "
    ws["F2"] = None                                    # колонка уточнений, не участник
    ws["F3"] = "clarificări"
    ws["G2"] = "Ofertant: Labromed Laborator SRL"
    ws["G3"] = "Producator: P.J. Dahlhausen & Co. GmbH\nȚara: Germania\nModel: 43.018.00.013"
    ws["G4"] = "DM000654009 Rg04-000259 09-11-2023"
    ws["H2"] = "Ofertant: "                            # пустая колонка шаблона
    wb.create_sheet("Lot x")["E2"] = "Ofertant: шаблон"
    return wb


class WorkbookTest(unittest.TestCase):

    def test_reads_participants_only(self):
        offers = read_offers(lot_workbook())
        self.assertEqual([o.supplier for o in offers], ['"ALMICOM-TECH" S.R.L.', "Labromed Laborator SRL"])
        self.assertEqual(offers[1].dm_numbers, ["DM000654009"])
        self.assertEqual(offers[0].quality_docs, ["CE", "ISO"])

    def test_report_sheet_does_not_touch_lot_cells(self):
        wb = lot_workbook()
        before = wb["lot 23"]["G4"].value
        results = [r for o in read_offers(wb) for r in check_offer(o, {"DM000654009": RECORD}, set())]
        write_report(wb, results)
        self.assertEqual(wb.sheetnames[0], SHEET)
        self.assertEqual(wb["lot 23"]["G4"].value, before)
        statuses = [wb[SHEET].cell(row=r, column=6).value for r in (2, 3)]
        self.assertEqual(sorted(statuses), sorted([CHECK, OK]))

    def test_opens_workbook_with_removed_images(self):
        """Книга, из которой удалили xl/media ради размера, всё равно открывается."""
        from openpyxl.drawing.image import Image
        from PIL import Image as PILImage
        with tempfile.TemporaryDirectory() as d:
            src, dst = Path(d) / "a.xlsx", Path(d) / "b.xlsx"
            wb = lot_workbook()
            buf = io.BytesIO()
            PILImage.new("RGB", (4, 4)).save(buf, format="PNG")
            wb["lot 23"].add_image(Image(io.BytesIO(buf.getvalue())), "J10")
            wb.save(src)
            with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
                for item in zin.infolist():
                    if not item.filename.startswith("xl/media/"):
                        zout.writestr(item, zin.read(item.filename))
            self.assertEqual(len(read_offers(load_workbook_tolerant(dst))), 2)


if __name__ == "__main__":
    unittest.main()
