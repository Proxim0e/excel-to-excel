import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from src import classify as cl
from src.manifest import Manifest

SIG = "Digitally signed by Ion Popescu Date: 2026.07.30 11:12:06 EEST Reason: MoldSign Signature Location: Moldova "
A22_HEADER = ("Specificaţii tehnice Cod CPV Denumirea bunurilor Modelul articolului Ţara de origine "
              "Producătorul Specificarea tehnică deplină solicitată")
A23_TABLE = ("Cod CPV Nr. lot Denumirea poziţiei Unitatea de măsură Cantitatea Preţ unitar (fără TVA) "
             "Preţ unitar (cu TVA) Suma (fără TVA) Suma (cu TVA) Termenul de livrare")
FILLER = " rand de tabel cu date despre produs si cantitati livrate" * 3


class RuleTests(unittest.TestCase):
    def test_anexa22_by_title(self):
        c = cl.classify([A22_HEADER + FILLER])
        self.assertEqual((c.doc_type, c.confidence), ("anexa22", 0.9))

    def test_anexa23_without_title_by_column_headers(self):
        self.assertEqual(cl.classify(["Data: 31.07.2026 Lot: 7 " + A23_TABLE + FILLER]).doc_type, "anexa23")

    def test_both_annexes_in_one_file(self):
        c = cl.classify([A22_HEADER + FILLER, A23_TABLE + FILLER])
        self.assertEqual(c.doc_type, "anexa22+23")
        self.assertEqual(c.contains, ["anexa22", "anexa23"])

    def test_table_continuation_inherits_type(self):
        c = cl.classify([A22_HEADER + FILLER, "12 Electrod ECG Fiab Italia" + FILLER])
        self.assertEqual(c.page_types, ["anexa22", "anexa22"])

    def test_signature_stamp_is_not_the_title(self):
        c = cl.classify([SIG + "Cerere de participare catre CAPCS" + FILLER])
        self.assertEqual((c.doc_type, c.confidence), ("cerere_participare", 0.9))

    def test_scan_with_only_signature_uses_filename(self):
        c = cl.classify([SIG, SIG], "lotul 149 reg.semnat.pdf")
        self.assertEqual((c.doc_type, c.confidence, c.has_text), ("amdm", 0.3, False))
        self.assertEqual(c.page_types, [None, None])

    def test_scan_without_hint_is_no_text(self):
        self.assertEqual(cl.classify([SIG], "scan001.pdf").doc_type, cl.NO_TEXT)

    def test_glued_filename(self):
        c = cl.classify([""], "7certificatprivindlipsarestantelor-0926657semnat.pdf")
        self.assertEqual(c.doc_type, "certificat_restante")

    def test_declaration_mentioning_specs_is_not_anexa22(self):
        letterhead = "LifeMed Group SRL MD-3733 mun. Chisinau codul fiscal 1016600012345 tel 022 " * 9
        page = letterhead + "Declaratie privind prezentarea specificatiilor tehnice si a mostrelor" + FILLER
        self.assertEqual(cl.classify([page]).doc_type, "declaratie")

    def test_bundle_lists_sections(self):
        pages = ["Anexa nr. 7 Cerere de participare" + FILLER, "Anexa nr. 8 Declaratie privind valabilitatea ofertei"
                 + FILLER, A22_HEADER + FILLER, A23_TABLE + FILLER]
        c = cl.classify(pages, "Oferta.signed.pdf")
        self.assertEqual(c.doc_type, "cerere_participare")
        self.assertEqual(c.contains, ["cerere_participare", "declaratie_valabilitate", "anexa22", "anexa23"])

    def test_iso_in_letterhead_is_not_a_certificate(self):
        letterhead = ("m-inter-farma s.a. idno 1003600005263 adresa: grenoble 21 chisinau tel: +373 22 904 005 "
                      "iso 9001:2015 web: www.minterfarma.md ") * 3
        page = letterhead + "catre centrul pentru achizitii " * 10 + "Declaratie privind perioada de garantie"
        self.assertEqual(cl.classify([page]).doc_type, "declaratie")

    def test_quality_certificate_without_iso_in_title(self):
        page = ("Certificate No. IT334924 certification body Bureau Veritas Italia certifies that the management "
                "system of the above organization has been audited" + FILLER)
        self.assertEqual(cl.classify([page]).doc_type, "iso")

    def test_sample_declaration_with_one_column_header_is_not_anexa22(self):
        letterhead = "SC Denolga Medical LTD Banulescu Bodoni str. 45 Chisinau tel 022 260 602 " * 9
        page = letterhead + "tabel: Nr. lot, modelul articolului, cantitatea mostrelor" + FILLER
        self.assertNotEqual(cl.classify([page]).doc_type, "anexa22")

    def test_scan_filename_with_standard_number(self):
        self.assertEqual(cl.classify([SIG], "lot 72 iso 13485.semnat.pdf").doc_type, "iso")

    def test_groups(self):
        self.assertEqual(cl.group_of("duae"), "Documente generale")
        self.assertEqual(cl.group_of("cerere_participare", "cerere_participare,anexa22,anexa23"), "anexa22/anexa23")
        self.assertEqual(cl._label_groups("Document tehnic/ISO/CE"), ["Document tehnic", "ISO/CE/DM"])
        self.assertEqual(cl._label_groups("ISO/CE"), ["ISO/CE/DM"])
        self.assertEqual(cl._label_groups("anexa22/anexa23"), ["anexa22/anexa23"])

    def test_catalog_page_by_product_markers(self):
        page = "upnmed ekg cable www.upnmed.com brand/compatible P/N description disposable 4.0 mm connector" * 2
        self.assertEqual(cl.classify([page]).doc_type, "datasheet")


class ReviewCycleTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        m = Manifest(self.root / "manifest.sqlite")
        self.texts = {}
        for i, (name, text) in enumerate([("anexa22.pdf", A22_HEADER + FILLER), ("misc.pdf", "zzz" + FILLER)]):
            p = self.root / "lot1" / "acme" / name
            p.parent.mkdir(parents=True, exist_ok=True)
            p.write_bytes(f"%PDF-{i}".encode())
            self.texts[str(p)] = [text]
            m.record(tender_id="T", lot_no=1, lot_title="Lot 1", participant="ACME", original_name=name,
                     url=f"u{i}", path=p, status="success")
        m.close()
        patcher = mock.patch.object(cl, "extract_pages", side_effect=lambda path: self.texts[str(path)])
        patcher.start()
        self.addCleanup(patcher.stop)

    def tearDown(self):
        self.tmp.cleanup()

    def types(self):
        conn = sqlite3.connect(self.root / "manifest.sqlite")
        try:
            return dict(conn.execute("SELECT d.doc_type, d.type_source FROM documents d"))
        finally:
            conn.close()

    def test_human_fix_survives_reclassification(self):
        self.assertEqual(sum(cl.classify_tender(self.root).values()), 2)
        self.assertEqual(cl.classify_tender(self.root), {})  # уже классифицированы

        xlsx = cl.export_review(self.root)
        wb = load_workbook(xlsx)
        ws = wb[cl.REVIEW_SHEET]
        header = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            if row[header.index("Тип (система)")].value == cl.OTHER:
                row[header.index(cl.REVIEW_COLUMN)].value = "datasheet"
        wb.save(xlsx)

        applied, errors = cl.apply_review(self.root)
        self.assertEqual((applied, errors), (1, []))
        cl.classify_tender(self.root, force=True)
        self.assertEqual(self.types(), {"anexa22": "rule", "datasheet": "human"})

    def label_all(self, labels_by_system_type):
        xlsx = cl.export_review(self.root)
        wb = load_workbook(xlsx)
        ws = wb[cl.REVIEW_SHEET]
        header = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            system_type = row[header.index("Тип (система)")].value
            row[header.index(cl.REVIEW_COLUMN)].value = labels_by_system_type[system_type]
        wb.save(xlsx)
        return xlsx

    def test_unknown_type_in_review_is_reported(self):
        cl.classify_tender(self.root)
        self.label_all({"anexa22": None, cl.OTHER: "nu exista"})
        applied, errors = cl.apply_review(self.root)
        self.assertEqual(applied, 0)
        self.assertIn("неизвестный тип", errors[0])

    def test_group_label_confirms_or_asks_for_type(self):
        cl.classify_tender(self.root)
        self.label_all({"anexa22": "anexa22", cl.OTHER: "Documente generale"})
        applied, errors = cl.apply_review(self.root)
        self.assertEqual(applied, 1)  # anexa22 - это и тип, и группа
        self.assertEqual(len(errors), 1)
        self.assertIn("выберите конкретный тип", errors[0])

    def test_evaluate_against_gold(self):
        cl.classify_tender(self.root)
        gold = self.label_all({"anexa22": "anexa22", cl.OTHER: "Documente generale"})
        r = cl.evaluate(self.root, gold)
        self.assertEqual((r["total"], r["correct"]), (2, 1))
        self.assertEqual(r["errors"][0][:3], ("Documente generale", cl.OTHER, "Document tehnic"))


if __name__ == "__main__":
    unittest.main()
