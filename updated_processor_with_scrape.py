#!/usr/bin/env python3
"""
Интегрированный скрипт с параллельным скрапингом лотов.

Изменение по вашей просьбе:
- При записи участников заполняется только колонка E (Ofertant, цена, формулы).
- Для каждого следующего участника создаётся новая колонка непосредственно перед колонкой
  "Note" (чтобы Note всегда оставалась последней). Новая колонка создаётся методом
  insert_cols и затем копируются стили/ширина из колонки E в новую колонку.
- Формулы для каждой новой колонки устанавливаются программно (=<col>8/A1 и =<col>8/D8).
- Если участников нет — E11 заполняется текстом "Achiziţia nu a avut loc" и оформляется.
- Если процент >= 130% — ячейка процента выделяется красным; если ВСЕ участники лота >=130%,
  вкладка помечается красным.
- Парсинг цен, параллелизм и остальная логика сохранены.

Примечание о копировании форматов: openpyxl поддерживает копирование стиля ячейки
за ячейкой — и мы копируем font, fill, border, number_format, alignment и protection.
Копирование merged-диапазонов для новых колонок не выполняется (можно добавить, но может
сложниться), обычно шаблон у вас настроен так, что колонки для участников — простые,
а крупные merged-области (заголовок, примечания) остаются нетронутыми.

Автор: Proxim0e (адаптация)
Дата: 2025-11-20
"""
import os
import re
from datetime import datetime
from urllib.parse import urljoin
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border
from openpyxl.cell import Cell

# Папка для ресурсов
RESOURCE_DIR = './resources'
TEMPLATE_FILE = os.path.join(RESOURCE_DIR, 'sample_model.xlsx')

# Ссылки и скрапинг
BASE = "https://achizitii.md"
TENDER_URL = "https://achizitii.md/ro/public/tender/21463176/"
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; scraper/1.0; +https://example.com/bot)"}
LOT_LINK_RE = re.compile(r'/ro/public/tender/\d+/lot/\d+/?$')

# Параметр параллелизма
MAX_WORKERS = 16  # можно настроить

# Названия столбцов/целевых ячеек
COLS = {
    'nr_lot':            ['Nr. Lot', 'Nr Lot'],
    'denumire':          ['Denumirea Lotului\n04.09.2025', 'Denumirea Lotului', 'Denumire Lot NEW2'],
    'specificatie':      ['Specificația Tehnică\n04.09.2025', 'Specificația Тehnică', 'Specificarea техническая NEW2'],
    'unitate_masura':    ['Unitatea de măsură'],
    'cantitate_total':   ['Cantitatea Totală'],
    'suma_alocata':      ['Suma alocată'],
}

TARGET_MAP = {
    "denumire": ("B1",),
    "specificatie": ("B11",),
    "unitate_masura": ("C9",),
    "cantitate_total": ("A1",),
    "suma_alocata": ("D8",)
}

# Цвета и шрифты
NOT_HELD_FILL_DARK = PatternFill(start_color="001f4d", end_color="001f4d", fill_type="solid")  # dark navy
NOT_HELD_FONT_WHITE = Font(color="FFFFFF", bold=True)

HIGH_PERCENT_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")  # red
HIGH_PERCENT_FONT = Font(color="000000", bold=True)

# threshold for percent (130% = 1.3)
PERCENT_THRESHOLD = 1.3

# -------------------
# Утилиты и парсинг
# -------------------
def find_file(prefix, ext='.xlsx'):
    for fname in os.listdir(RESOURCE_DIR):
        if fname.startswith(prefix) and fname.endswith(ext):
            return os.path.join(RESOURCE_DIR, fname)
    raise FileNotFoundError(f"Файл с префиксом '{prefix}' не найден в директории {RESOURCE_DIR}!")


def safe_col_idx(names, headers):
    if isinstance(names, str):
        names = [names]
    for name in names:
        try:
            return headers.index(name) + 1
        except ValueError:
            continue
    return None


def safe_row_get(row, keys, headers):
    idx = safe_col_idx(keys, headers)
    if idx is not None:
        return row[idx - 1]
    return None


def get_soup(url):
    try:
        resp = requests.get(url, headers=HEADERS, timeout=20)
        resp.raise_for_status()
        return BeautifulSoup(resp.text, "html.parser")
    except Exception as e:
        print(f"[ERROR] Ошибка загрузки {url}: {e}")
        return None


def extract_lot_links(tender_url):
    soup = get_soup(tender_url)
    if not soup:
        return []
    links = set()
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        full = urljoin(BASE, href)
        if LOT_LINK_RE.search(href) or LOT_LINK_RE.search(full.replace(BASE, "")):
            links.add(full.split('#')[0].rstrip("/"))
    return sorted(links)


def strip_price_label(s: str) -> str:
    if not s:
        return s
    s = s.strip()
    s = re.sub(r'(?i)prețul ofertei[:\s]*', '', s)
    s = re.sub(r'(?i)preţul ofertei[:\s]*', '', s)
    s = re.sub(r'(?i)preţul[:\s]*', '', s)
    s = re.sub(r'(?i)preț[:\s]*', '', s)
    s = re.sub(r'(?i)preţ[:\s]*', '', s)
    s = re.sub(r'(?i)pre[uț]ț?[:\s]*', '', s)
    return s.strip()


def parse_lot_html(soup):
    if not soup:
        return None, []

    title = None
    for sel in ("h1", ".tender__page__title h1", ".tender__page__title"):
        el = soup.select_one(sel)
        if el and el.get_text(strip=True):
            title = el.get_text(strip=True)
            break
    if not title:
        mt = soup.find("title")
        title = mt.get_text(strip=True) if mt else "(untitled)"

    participants = []
    infos = soup.select(".participant-container-body-info")
    if infos:
        for info in infos:
            name = ""
            nnode = info.select_one(".participant-container-body-item.participant-title, .participant-title, .participant-title-div")
            if nnode:
                name = nnode.get_text(separator=" ", strip=True)
            price = ""
            pnode = info.select_one(".participant-container-body-item.participant-price, .participant-price")
            if pnode:
                price_text = pnode.get_text(separator=" ", strip=True)
                price = strip_price_label(price_text)
            if name or price:
                participants.append((name, price))
    else:
        title_nodes = soup.select(".participant-container-body .participant-title, .participant-title")
        price_nodes = soup.select(".participant-container-body .participant-price, .participant-price")
        if title_nodes and price_nodes and len(title_nodes) == len(price_nodes):
            for t, p in zip(title_nodes, price_nodes):
                price_clean = strip_price_label(p.get_text(strip=True))
                participants.append((t.get_text(strip=True), price_clean))

    return title, participants


def parse_lot_page(lot_url):
    try:
        soup = get_soup(lot_url)
        if not soup:
            return {'url': lot_url, 'title': None, 'participants': [], 'error': 'load_failed'}
        title, participants = parse_lot_html(soup)
        return {'url': lot_url, 'title': title, 'participants': participants, 'error': None}
    except Exception as e:
        return {'url': lot_url, 'title': None, 'participants': [], 'error': str(e)}


def clean_company_name(raw_name):
    if not raw_name:
        return raw_name
    s = raw_name.strip()
    s = re.sub(r'(?i)denumirea participantului[:\s]*', '', s).strip()
    s = re.sub(r'^\s*Denumirea[:\s]*', '', s, flags=re.I).strip()
    s = re.split(r'\s+Preţ|Preţul|Preț', s)[0].strip()
    return s


def parse_price_to_number(price_str):
    if not price_str:
        return None
    s = price_str.strip()
    m = re.search(r'[\d\s\.,]+', s)
    if not m:
        return None
    s = m.group(0).strip()
    s = s.replace('\u00A0', ' ')
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '')
            s = s.replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        s = s.replace(' ', '')
        s = s.replace(',', '.')
    else:
        s = s.replace(' ', '')
    s = re.sub(r'[^\d\.]', '', s)
    if not s:
        return None
    try:
        if '.' in s:
            return float(s)
        else:
            return int(s)
    except Exception:
        return None


def clean_denumire_for_B1(denumire):
    if not denumire:
        return ""
    s = str(denumire).strip()
    s = re.sub(r'(?i)^\s*Lot(?:ul)?\s*nr\.?\s*\d+\s*[-–:\)]*\s*', '', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s


def normalize_spaces(s):
    if not s:
        return s
    return re.sub(r'\s+', ' ', str(s)).strip()


# -------------------
# Helpers for column copying
# -------------------
def copy_column_styles(ws, src_col_idx, dst_col_idx, max_row):
    """
    Копирует стиль (font, fill, border, number_format, alignment, protection) и ширину из src -> dst.
    Не копирует значения.
    """
    src_letter = get_column_letter(src_col_idx)
    dst_letter = get_column_letter(dst_col_idx)
    # width
    try:
        if ws.column_dimensions.get(src_letter) and ws.column_dimensions.get(src_letter).width:
            ws.column_dimensions[dst_letter].width = ws.column_dimensions[src_letter].width
    except Exception:
        pass

    for r in range(1, max_row + 1):
        src_cell = ws.cell(row=r, column=src_col_idx)
        dst_cell = ws.cell(row=r, column=dst_col_idx)
        # Copy style attributes
        try:
            dst_cell.font = src_cell.font.copy()
        except Exception:
            pass
        try:
            dst_cell.fill = src_cell.fill.copy()
        except Exception:
            pass
        try:
            dst_cell.border = src_cell.border.copy()
        except Exception:
            pass
        try:
            dst_cell.number_format = src_cell.number_format
        except Exception:
            pass
        try:
            dst_cell.alignment = src_cell.alignment.copy()
        except Exception:
            pass
        try:
            dst_cell.protection = src_cell.protection.copy()
        except Exception:
            pass


# -------------------
# Основной рабочий код
# -------------------
def main():
    try:
        parent_file = find_file('-')
        print(f"Родительский файл: {parent_file}")
        print(f"Шаблонный файл: {TEMPLATE_FILE}")

        wb_parent = load_workbook(parent_file, data_only=True)
        ws_parent = wb_parent.active
        wb_child = load_workbook(TEMPLATE_FILE)
        sheet_names = wb_child.sheetnames
        template_sheet = wb_child[sheet_names[0]]

        # Инициализация mapping листов (lot N -> Worksheet)
        lot_regex = re.compile(r'^lot (\d+)$')
        parent_lot_sheets = {}
        for sname in wb_child.sheetnames:
            m = lot_regex.match(sname)
            if m:
                try:
                    num = int(m.group(1))
                    parent_lot_sheets[num] = wb_child[sname]
                except Exception:
                    continue
        print(f"Найдено в шаблоне {len(parent_lot_sheets)} готовых листов.")

        # Заголовки родительского файла
        headers = [cell.value for cell in ws_parent[1]]
        print("Заголовки в родительском файле:", headers)

        # Определим индекс колонки Note в шаблоне (ищем заголовок "Note" в первой строк)
        def find_note_col(ws):
            for c, cell in enumerate(ws[1], start=1):
                if cell.value and isinstance(cell.value, str) and cell.value.strip().lower() == "note":
                    return c
            # fallback: если не найден — возвращаем последнюю существующую колонку +1 (т.е. Note будет добавлена в конец)
            return ws.max_column + 1

        note_col_template = find_note_col(template_sheet)
        print(f"[INFO] Note column in template: {get_column_letter(note_col_template)} (index {note_col_template})")

        # Создание/обновление листов по родительскому файлу
        for row in ws_parent.iter_rows(min_row=2, values_only=True):
            nr_lot = safe_row_get(row, COLS['nr_lot'], headers)
            if nr_lot is None or (isinstance(nr_lot, str) and not nr_lot.strip()):
                continue
            try:
                nr_lot_int = int(float(nr_lot))
            except Exception:
                print(f"[WARN] Некорректный номер лота: {nr_lot}")
                continue

            denumire = safe_row_get(row, COLS['denumire'], headers)
            specificatie_data = safe_row_get(row, COLS['specificatie'], headers)

            if nr_lot_int in parent_lot_sheets:
                ws_target = parent_lot_sheets[nr_lot_int]
                if specificatie_data:
                    existing = ws_target["B11"].value or ""
                    add = str(specificatie_data).strip()
                    if add and add not in existing:
                        new_val = (existing + "\n" + add).strip() if existing else add
                        ws_target["B11"] = new_val
                        print(f"[Excel] Дополнена спецификация для lot {nr_lot_int}")
                continue

            new_sheet = wb_child.copy_worksheet(template_sheet)
            new_title = f"lot {nr_lot_int}"
            new_sheet.title = new_title
            parent_lot_sheets[nr_lot_int] = new_sheet
            print(f"[Excel] Создаем лист: {new_sheet.title}")

            # Записываем denumire в target cell (как делаем для B11) — принудительно "Lot nr. N {denumire}"
            den_text_clean = clean_denumire_for_B1(denumire)
            value_to_write = normalize_spaces(f"Lot nr. {nr_lot_int} {den_text_clean}".strip())
            for cell in TARGET_MAP['denumire']:
                new_sheet[cell] = value_to_write
                try:
                    new_sheet[cell].font = Font(bold=True)
                except Exception:
                    pass
            print(f"[EXCEL] Установлен заголовок для листа '{new_sheet.title}' в {TARGET_MAP['denumire']}: {value_to_write!r}")

            # Заполнение остальных целей (кроме denumire)
            for key, target_cells in TARGET_MAP.items():
                if key == "denumire":
                    continue
                if key == "specificatie":
                    continue
                value = safe_row_get(row, COLS[key], headers)
                if value is not None:
                    for cell in target_cells:
                        new_sheet[cell] = value

            # initial specificatie
            if specificatie_data:
                new_sheet["B11"] = str(specificatie_data).strip()

        # -------------------
        # Параллельный скрапинг лотов
        # -------------------
        print("[SCRAPE] Собираем ссылки на лоты...")
        lot_links = extract_lot_links(TENDER_URL)
        print(f"[SCRAPE] Найдено {len(lot_links)} ссылок на лоты. Параллельность: {MAX_WORKERS} воркеров.")

        results = []
        if lot_links:
            workers = min(MAX_WORKERS, max(2, len(lot_links)))
            with ThreadPoolExecutor(max_workers=workers) as ex:
                future_to_url = {ex.submit(parse_lot_page, url): url for url in lot_links}
                for fut in as_completed(future_to_url):
                    url = future_to_url[fut]
                    try:
                        res = fut.result()
                    except Exception as e:
                        print(f"[ERROR] Ошибка при обработке {url}: {e}")
                        res = {'url': url, 'title': None, 'participants': [], 'error': str(e)}
                    results.append(res)
                    if res.get('error'):
                        print(f"[SCRAPE] {url} -> error: {res['error']}")
                    else:
                        tcount = len(res.get('participants') or [])
                        title_sample = normalize_spaces(res.get('title') or "")
                        print(f"[SCRAPE] {url} -> title: {title_sample!r}, participants: {tcount}")

        # -------------------
        # Применяем результаты к Excel (в главном потоке)
        # -------------------
        for r in results:
            lot_url = r.get('url')
            title = r.get('title')
            participants = r.get('participants') or []
            lot_number = None
            if title:
                m = re.search(r'Lot(?:ul)?\s*nr\.?\s*\.*\s*(\d+)', title, flags=re.I)
                if m:
                    lot_number = int(m.group(1))
            if lot_number is None:
                print(f"[WARN] Не удалось извлечь номер лота из title для URL {lot_url!r} title={title!r}. Пропускаю.")
                continue

            if lot_number not in parent_lot_sheets:
                print(f"[WARN] Для lot {lot_number} нет листа в книге (пропуск обновления участников).")
                continue

            ws = parent_lot_sheets[lot_number]

            # Find current note column for this worksheet (scan first row)
            def find_note_col_ws(ws_loc):
                for c, cell in enumerate(ws_loc[1], start=1):
                    if cell.value and isinstance(cell.value, str) and cell.value.strip().lower() == "note":
                        return c
                return ws_loc.max_column + 1

            note_col = find_note_col_ws(ws)
            # Ensure note_col is at least >5; if note is before E for some reason, set to E+1
            if note_col <= 5:
                note_col = 6

            if not participants:
                # вставляем фиксированный текст + оформление (тёмно-синий фон и белый текст)
                print(f"[EXCEL] lot {lot_number}: участников не найдено — вставляем 'Achiziţia nu a avut loc' в E11.")
                ws["E11"] = "Achiziţia nu a avut loc"
                ws["E11"].fill = NOT_HELD_FILL_DARK
                ws["E11"].alignment = Alignment(horizontal="center", vertical="center")
                try:
                    ws["E11"].font = NOT_HELD_FONT_WHITE
                except Exception:
                    pass
                # помечаем вкладку листа тем же цветом
                try:
                    ws.sheet_properties.tabColor = "001f4d"
                except Exception:
                    pass
                continue

            # Если есть участники: заполняем E для первого, остальные создаём копией оформления колонки E и вставляем перед Note
            max_row = ws.max_row if ws.max_row > 1 else 50  # безопасный максимум для копирования стилей
            src_col = 5  # E
            high_flags = []

            for idx, (raw_name, raw_price) in enumerate(participants, start=1):
                # For idx==1 -> write directly into E
                if idx == 1:
                    target_col = src_col
                else:
                    # insert new column immediately before current note_col
                    insert_at = note_col  # insert at this index, note will shift right
                    ws.insert_cols(insert_at, amount=1)
                    target_col = insert_at
                    # copy styles from src_col (E) into new column (target_col)
                    copy_column_styles(ws, src_col, target_col, max_row)
                    # after insertion note_col increases by 1
                    note_col += 1

                col_letter = get_column_letter(target_col)
                name = clean_company_name(raw_name) or "(без имени)"
                price_value = parse_price_to_number(raw_price) if raw_price else None

                # cells
                cell_name = f"{col_letter}2"
                cell_price = f"{col_letter}8"
                cell_formula_divA1 = f"{col_letter}7"
                cell_formula_divD8 = f"{col_letter}9"

                # write ofertant
                ws[cell_name] = f"Ofertant: {name}"
                try:
                    ws[cell_name].font = Font(bold=True)
                except Exception:
                    pass

                # write price (only numeric)
                if price_value is not None:
                    ws[cell_price] = price_value
                else:
                    ws[cell_price] = ""

                # formulas
                ws[cell_formula_divA1] = f"={col_letter}8/A1"
                ws[cell_formula_divD8] = f"={col_letter}8/D8"

                # evaluate percent with D8 if possible and set red style if >= threshold
                is_high = False
                try:
                    d8_raw = ws["D8"].value
                    d8_val = None
                    if isinstance(d8_raw, (int, float)):
                        d8_val = float(d8_raw)
                    else:
                        d8_val = parse_price_to_number(str(d8_raw)) if d8_raw is not None else None
                    if price_value is not None and d8_val not in (None, 0):
                        percent = float(price_value) / float(d8_val)
                        if percent >= PERCENT_THRESHOLD:
                            is_high = True
                            ws[cell_formula_divD8].fill = HIGH_PERCENT_FILL
                            ws[cell_formula_divD8].font = HIGH_PERCENT_FONT
                            ws[cell_formula_divD8].alignment = Alignment(horizontal="center", vertical="center")
                except Exception:
                    pass

                high_flags.append(is_high)
                print(f"[EXCEL] lot {lot_number}: записан участник в {col_letter} -> '{name}' / price={price_value} high={is_high}")

            # после обработки всех участников: если есть участники и ВСЕ high_flags=True => пометить вкладку красным
            try:
                if high_flags and all(high_flags):
                    ws.sheet_properties.tabColor = "FF0000"
                    print(f"[EXCEL] lot {lot_number}: все участники >={int(PERCENT_THRESHOLD*100)}% -> вкладка помечена красным")
            except Exception:
                pass

        # Сохранение файла
        parent_basename = os.path.basename(parent_file).replace('-', '')
        timestamp = datetime.now().strftime("%d.%m.%Y_%H-%M-%S")
        output_name = f"{RESOURCE_DIR}/{os.path.splitext(parent_basename)[0]}_{timestamp}_upd.xlsx"
        wb_child.save(output_name)
        print(f"\nГотово! Файл сохранён как {output_name}")

    except FileNotFoundError as e:
        print(e)
    except Exception as e:
        print(f"Неизвестная ошибка: {e}")


if __name__ == "__main__":
    main()