import re
import copy
from typing import Optional, List, Dict
from pathlib import Path
import logging

logger = logging.getLogger(__name__)
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.config import (
    TEMPLATE_FILE, TARGET_MAP, COLS,
    APPEND_TO_EXISTING, PERCENT_THRESHOLD,
    NOT_HELD_FILL_DARK, NOT_HELD_FONT_WHITE
)
from src.models import LotData
from src.utils.text_helpers import prepare_lot_title_for_b1, normalize_spaces
from src.excel.styles import (
    NOT_HELD_FILL, NOT_HELD_FONT, NOT_HELD_ALIGNMENT,
    HIGH_PERCENT_FILL, HIGH_PERCENT_FONT, HIGH_PERCENT_ALIGNMENT,
    HEADER_FONT_BOLD
)


class ExcelManager:
    def __init__(self):
        # Загружаем книгу шаблона
        self.wb = load_workbook(TEMPLATE_FILE)
        self.template_sheet = self.wb[self.wb.sheetnames[0]]

        # Карта существующих листов (lot_number -> worksheet)
        self.lot_sheets: Dict[int, Worksheet] = {}
        self._scan_existing_sheets()

    def _scan_existing_sheets(self):
        """Сканирует книгу и ищет листы с названием 'lot N'."""
        lot_regex = re.compile(r'^lot (\d+)$')
        for sname in self.wb.sheetnames:
            m = lot_regex.match(sname)
            if m:
                try:
                    num = int(m.group(1))
                    self.lot_sheets[num] = self.wb[sname]
                except ValueError:
                    continue

        # Вычисляем максимальный номер уже существующего лота
        self.max_existing: int = max(self.lot_sheets.keys()) if self.lot_sheets else 0

    def get_or_create_sheet(self, lot_number: int) -> Optional[Worksheet]:
        """
        Возвращает объект листа для лота. Если листа нет — создает его.
        """
        nr = lot_number

        # Если лист уже есть
        if nr in self.lot_sheets:
            if not APPEND_TO_EXISTING and nr <= self.max_existing:
                return None
            return self.lot_sheets[nr]

        # Создаем новый лист
        new_sheet = self.wb.copy_worksheet(self.template_sheet)
        new_title = f"lot {nr}"
        new_sheet.title = new_title
        self.lot_sheets[nr] = new_sheet

        # Обратите внимание: Заполнение ячеек (B1, D8 и т.д.) мы перенесли в main.py
        return new_sheet

    @staticmethod
    def _fill_lot_metadata(ws, lot_data: LotData):
        """Заполняет заголовок и базовые поля лота из родительского файла (если данные переданы)"""
        # Внимание: здесь мы пока используем данные из объекта LotData,
        # но логически мета-дата (название, спецификация) приходит из родительского Excel.
        # В данном упрощении мы предполагаем, что LotData уже содержит всё нужное для заполнения
        # или мы вызываем этот метод только когда точно знаем, что есть данные.

        # Для примера, запишем заголовок, если он есть в данных
        if lot_data.title:
            clean_title = prepare_lot_title_for_b1(lot_data.title)
            val = f"Lot nr. {lot_data.number} {clean_title}".strip()

            # Запись в B1 (TARGET_MAP['denumire'])
            for cell_addr in TARGET_MAP.get("denumire", []):
                ws[cell_addr] = val
                ws[cell_addr].font = HEADER_FONT_BOLD

    @staticmethod
    def find_note_col_idx(ws) -> int:
        """Находит индекс колонки 'Note' или возвращает текущую последнюю + 1"""
        # Смотрим первую строку на наличие "note"
        for c, cell in enumerate(ws[1], start=1):
            if cell.value and isinstance(cell.value, str) and cell.value.strip().lower() == "note":
                return c
        return ws.max_column + 1

    @staticmethod
    def copy_column_styles(ws, src_col_idx: int, dst_col_idx: int, max_row: int):
        """Копирует стили из одной колонки в другую"""
        # 1. Копируем ШИРИНУ КОЛОНКИ (это свойство всей колонки)
        src_letter = get_column_letter(src_col_idx)
        dst_letter = get_column_letter(dst_col_idx)

        try:
            src_dim = ws.column_dimensions.get(src_letter)
            # Если в шаблоне у Е задана ширина, ставим её и новой колонке
            if src_dim and src_dim.width is not None:
                ws.column_dimensions[dst_letter].width = src_dim.width
                # Логируем: ширина УСТАНОВЛЕНА успешно
                logger.debug(f"Ширина колонки скопирована: {src_letter} ({src_dim.width}) -> {dst_letter}")
        except AttributeError as e:
            logger.error(f"Критическая ошибка при чтении ширины колонки {src_letter}: {e}")

        for r in range(1, max_row + 1):
            src_cell = ws.cell(row=r, column=src_col_idx)
            dst_cell = ws.cell(row=r, column=dst_col_idx)

            try:
                dst_cell.font = copy.copy(src_cell.font)
                dst_cell.fill = copy.copy(src_cell.fill)
                dst_cell.border = copy.copy(src_cell.border)
                dst_cell.number_format = src_cell.number_format
                dst_cell.alignment = copy.copy(src_cell.alignment)
                dst_cell.protection = copy.copy(src_cell.protection)

                # (Опционально) Раскомментируйте строку ниже, если хотите видеть лог для КАЖДОЙ строки.
                # Внимание: при больших файлах это создаст очень много сообщений в консоли.
                #logger.debug(f"Стили успешно скопированы для ячейки {dst_letter}{r}")

            except AttributeError as e:
                # ВАЖНО: Раньше здесь был 'pass', и ошибки игнорировались.
                # Теперь мы записываем, в какой именно ячейке не скопировался стиль.
                logger.debug(f"Не удалось скопировать стиль для строки {r} ({src_letter}{r}): {e}")

    def write_participants(self, lot_data: LotData):
        """
        Основная логика записи участников.
        Создает колонки, формулы и применяет стили.
        """
        ws = self.lot_sheets.get(lot_data.number)
        if not ws:
            # Если листа нет (например, мы пропустили его в get_or_create_sheet из-за APPEND_TO_EXISTING)
            if not APPEND_TO_EXISTING and lot_data.number <= self.max_existing:
                return
            else:
                # Если что-то пошло не так и листа нет, создадим его (для новых лотов)
                ws = self.get_or_create_sheet(lot_data.number)
                if not ws: return

        participants = lot_data.participants
        note_col_idx = self.find_note_col_idx(ws)
        if note_col_idx <= 5:
            note_col_idx = 6

        # Если участников нет
        if not participants:
            ws["E11"] = "Achiziţia nu a avut loc"
            ws["E11"].fill = NOT_HELD_FILL
            ws["E11"].font = NOT_HELD_FONT
            ws["E11"].alignment = NOT_HELD_ALIGNMENT
            try:
                ws.sheet_properties.tabColor = NOT_HELD_FILL_DARK
            except AttributeError:
                # Если вдруг свойство отсутствует, просто игнорируем
                pass
            logger.info(
                f"[EXCEL] Лот {lot_data.number}: участников не найдено — вставляем 'Achiziţia nu a avut loc' в E11.")
            return

        # Записываем участников
        max_row = ws.max_row if ws.max_row > 1 else 50
        src_col = 5  # Колонка E
        high_flags = []

        for idx, part in enumerate(participants, start=1):
            if idx == 1:
                target_col = src_col
            else:
                # Вставляем новую колонку перед Note
                ws.insert_cols(note_col_idx, amount=1)
                target_col = note_col_idx
                self.copy_column_styles(ws, src_col, target_col, max_row)
                note_col_idx += 1  # Note сдвинулась вправо

            col_letter = get_column_letter(target_col)

            # Пишем имя участника
            name_cell = f"{col_letter}2"
            ws[name_cell] = f"Ofertant: {part.name}"

            # Пишем цену
            price_cell = f"{col_letter}8"
            ws[price_cell] = part.price_val if part.price_val is not None else ""

            # Формулы
            # 7 строка: цена / A1
            ws[f"{col_letter}7"] = f"={col_letter}8/A1"
            # 9 строка: цена / D8
            percent_cell = f"{col_letter}9"
            ws[percent_cell] = f"={col_letter}8/D8"

            # Проверка высокого процента
            is_high = False
            try:
                d8_raw = ws["D8"].value
                d8_val = None
                if isinstance(d8_raw, (int, float)):
                    d8_val = float(d8_raw)

                # Конкретные ошибки деления и типов
                if part.price_val is not None and d8_val not in (None, 0):
                    percent = float(part.price_val) / float(d8_val)
                    if percent >= PERCENT_THRESHOLD:
                        is_high = True
                        ws[percent_cell].fill = HIGH_PERCENT_FILL
                        ws[percent_cell].font = HIGH_PERCENT_FONT
                        ws[percent_cell].alignment = HIGH_PERCENT_ALIGNMENT
            except (ZeroDivisionError, ValueError, TypeError):
                # Ловим ошибки деления на ноль или несовпадения типов, просто пропускаем подсветку
                pass

            high_flags.append(is_high)
            # ЛОГ: Записан участник
            logger.info(
                f"[EXCEL] Лот {lot_data.number}: записан участник в {col_letter} -> '{part.name}' / price={part.price_val} high={is_high}")

        # Если все high -> красная вкладка
        if high_flags and all(high_flags):
            try:
                ws.sheet_properties.tabColor = HIGH_PERCENT_FILL.start_color.rgb
            except AttributeError:
                pass
            # ЛОГ: Вкладка красная
            logger.info(
                f"[EXCEL] Лот {lot_data.number}: все участники >={int(PERCENT_THRESHOLD * 100)}% -> вкладка помечена красным")

    def save(self, output_path: str):
        """Сохраняет книгу по указанному пути"""
        self.wb.save(output_path)
        print(f"[EXCEL] File saved to {output_path}")