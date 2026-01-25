import re
from typing import Dict, List, Optional, Any
from pathlib import Path
from openpyxl import load_workbook

from src.config import RESOURCE_DIR, COLS


def find_parent_file() -> Path:
    """
    Ищет в папке resources файл, начинающийся с '-'.
    Выбрасывает FileNotFoundError, если не найден.
    """
    if not RESOURCE_DIR.exists():
        raise FileNotFoundError(f"Directory {RESOURCE_DIR} does not exist")

    for fname in RESOURCE_DIR.glob('-*'):
        if fname.suffix.lower() in ('.xlsx', '.xls'):
            return fname

    raise FileNotFoundError(f"Parent file starting with '-' not found in {RESOURCE_DIR}")


def get_parent_data() -> Dict[int, Dict[str, Any]]:
    """
    Читает родительский файл и возвращает данные по лотам.

    Returns:
        Словарь { lot_number: { key: value, ... } }
    """
    parent_path = find_parent_file()
    wb = load_workbook(parent_path, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]

    # Находим индексы нужных колонок по маппингу COLS
    col_indices = {}
    for key, possible_names in COLS.items():
        idx = _find_col_index(headers, possible_names)
        if idx:
            col_indices[key] = idx

    # Если не нашли номер лота, работать не можем
    if 'nr_lot' not in col_indices:
        raise ValueError("Column 'Nr. Lot' not found in parent file headers")

    lot_data = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        nr_lot_val = row[col_indices['nr_lot'] - 1]

        # Пропускаем пустые строки
        if nr_lot_val is None or (isinstance(nr_lot_val, str) and not nr_lot_val.strip()):
            continue

        try:
            # Номер лота может быть float (1.0) или string ("1")
            nr_lot = int(float(nr_lot_val))
        except (ValueError, TypeError):
            continue

        # Собираем данные
        data = {}
        for key in ['denumire', 'specificatie', 'unitate_masura', 'cantitate_total', 'suma_alocata']:
            if key in col_indices:
                data[key] = row[col_indices[key] - 1]
            else:
                data[key] = None

        lot_data[nr_lot] = data

    return lot_data


def get_tender_url_from_parent() -> Optional[str]:
    """
    Извлекает Tender URL из родительского файла.
    Ищет колонку с заголовком, содержащим 'site'/'url'/'link'.
    """
    try:
        parent_path = find_parent_file()
    except FileNotFoundError:
        return None

    wb = load_workbook(parent_path, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    # Ищем колонку с URL
    url_col_idx = None
    for i, h in enumerate(headers):
        if not h:
            continue
        h_lower = str(h).lower()
        if h_lower in ("site", "url", "link") or any(k in h_lower for k in ("site", "url", "link")):
            url_col_idx = i + 1
            break

    if url_col_idx:
        # Ищем первое непустое значение
        for row in ws.iter_rows(min_row=2, values_only=True):
            val = row[url_col_idx - 1]
            if val and isinstance(val, str) and val.strip().lower().startswith("http"):
                return val.strip()

    return None


def _find_col_index(headers: List[Optional[str]], possible_names: List[str]) -> Optional[int]:
    """Вспомогательная функция: ищет индекс колонки по списку возможных имен."""
    for name in possible_names:
        # Чувствительность к регистру не важна, но сохраняем порядок имен в списке
        if name is None:
            continue
        for i, h in enumerate(headers):
            if h is None:
                continue
            if str(h).strip().lower() == str(name).strip().lower():
                return i + 1
    return None