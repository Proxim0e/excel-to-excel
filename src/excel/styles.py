from openpyxl.styles import Font, PatternFill, Alignment

# ---------------------------------------------------------------------------
# Стили для ситуации "Торги не состоялись" (Achiziţia nu a avut loc)
# Тёмно-синяя заливка, белый жирный шрифт
# ---------------------------------------------------------------------------
NOT_HELD_FILL = PatternFill(start_color="001f4d", end_color="001f4d", fill_type="solid")
NOT_HELD_FONT = Font(color="FFFFFF", bold=True)
NOT_HELD_ALIGNMENT = Alignment(horizontal="center", vertical="center")

# ---------------------------------------------------------------------------
# Стили для выделения высокого процента (High Percent)
# Красная заливка, черный шрифт
# ---------------------------------------------------------------------------
HIGH_PERCENT_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
HIGH_PERCENT_FONT = Font(color="000000", bold=True)
HIGH_PERCENT_ALIGNMENT = Alignment(horizontal="center", vertical="center")

# ---------------------------------------------------------------------------
# Общие стили
# ---------------------------------------------------------------------------
HEADER_FONT_BOLD = Font(bold=True)