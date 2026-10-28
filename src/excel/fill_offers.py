"""
Этап 3: запись распознанного из Anexa 22 в листы лотов result-файла.

Для каждого участника (колонка от E, «Ofertant: <имя>» в строке 2):
  строка 3  - «Producator: X / Țara: Y / Model: Z» (в шаблоне уже есть подписи без значений);
  строка 4  - номера DM, а если их нет - «Standarde» из Anexa 22 (CE/ISO...);
  строка 11 - предложенная (декларативная) спецификация,
              плюс уровень 2: // "цитаты из даташита", p[страница], документ;
              нет технических документов - //docs, не найден параметр - // <параметр>.
Ячейка строки 11 красится по итогу уровня 2 (жёлтый - нужны уточнения, зелёный - подтверждено,
светло-зелёный - только декларативно), вкладка листа - по худшему участнику.
Пишется только в пустые ячейки (или в шаблон с пустыми подписями) - ничего не затирается;
уже окрашенные вкладки (например, «не состоялась») не перекрашиваются.

Запуск отдельно:  python -m src.excel.fill_offers <result.xlsx> <папка тендера> [--out <файл>] [--no-evidence]
"""
import argparse
import logging
import re
import sqlite3
import sys
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill

from src.anexa22.evaluate import participant_key
from src.anexa22.evidence import GREEN, LIGHT_GREEN, YELLOW, check_offer
from src.anexa22.gold import parse_row3
from src.manifest import MANIFEST_NAME

logger = logging.getLogger(__name__)

FIRST_PARTICIPANT_COL = 5   # E
ROW_NAME, ROW_PRODUCER, ROW_DM, ROW_SPEC = 2, 3, 4, 11
SHEET_RE = re.compile(r"lot\s*(\d+)", re.IGNORECASE)
EXCEL_CELL_LIMIT = 32767
WRAP = Alignment(wrap_text=True, vertical="top")

# Цвета как в Excel Ivan (21641880 CDM27 B): зелёный 00B050, жёлтый FFFF00; светло-зелёный -
# стандартный светло-зелёный Excel (в образце не встречается)
COLORS = {GREEN: "FF00B050", LIGHT_GREEN: "FF92D050", YELLOW: "FFFFFF00"}
# Вкладка листа: хоть один участник требует уточнений - жёлтая; все подтверждены - зелёная.
# Шаблон уже красит вкладки в зелёный, поэтому сохраняем только осмысленные цвета («не состоялась»).
TAB_PRIORITY = (YELLOW, LIGHT_GREEN, GREEN)
KEEP_TAB_COLORS = {"FF001F4D", "001F4D"}

OFFER_FIELDS = ("model", "country", "producer", "spec_offered", "standards", "dm")


def load_offers(conn: sqlite3.Connection) -> Dict[Tuple[int, str], dict]:
    """(лот, ключ участника) -> поля; несколько строк одного лота склеиваются (разные значения через \\n)."""
    grouped: Dict[Tuple[int, str], dict] = {}
    rows = conn.execute(f"SELECT lot_no, participant, {', '.join(OFFER_FIELDS)} FROM anexa22_offers "
                        "WHERE lot_no IS NOT NULL ORDER BY id")
    for lot, participant, *values in rows:
        rec = grouped.setdefault((lot, participant_key(participant)),
                                 {"participant": participant, **{f: [] for f in OFFER_FIELDS}})
        for name, value in zip(OFFER_FIELDS, values):
            value = (value or "").strip()
            if value and value not in rec[name]:
                rec[name].append(value)
    return grouped


def producer_cell(rec: dict) -> str:
    return (f"Producator: {'; '.join(rec['producer'])}\n"
            f"Țara: {'; '.join(rec['country'])}\n"
            f"Model: {'; '.join(rec['model'])}")


def dm_cell(rec: dict) -> str:
    dms = [d for value in rec["dm"] for d in value.split(",") if d]
    if dms:
        return "\n".join(dict.fromkeys(dms))
    return "\n".join(rec["standards"])


def _is_empty(value, row: int) -> bool:
    if value is None or not str(value).strip():
        return True
    # Шаблон строки 3: подписи «Producator: / Țara: / Model:» без значений
    return row == ROW_PRODUCER and not any(parse_row3(value))


def _write(ws, row: int, col: int, value: str) -> bool:
    cell = ws.cell(row=row, column=col)
    if not value or not _is_empty(cell.value, row):
        return False
    cell.value = value[:EXCEL_CELL_LIMIT]
    cell.alignment = WRAP
    return True


def _fill(status: str) -> PatternFill:
    return PatternFill("solid", fgColor=COLORS[status])


def fill_result(result_path: Path, tender_dir: Path, out_path: Optional[Path] = None,
                evidence: bool = True) -> dict:
    """Заполняет строки 3, 4, 11 листов лотов. Возвращает статистику."""
    manifest = Path(tender_dir) / MANIFEST_NAME
    if not manifest.is_file() or manifest.stat().st_size == 0:
        # sqlite3.connect молча создал бы пустой файл на месте отсутствующего манифеста
        raise FileNotFoundError(f"Нет манифеста тендера: {manifest}. Сначала запустите скачивание (main.py).")
    conn = sqlite3.connect(manifest)
    offers = load_offers(conn)
    wb = load_workbook(result_path)
    stats = defaultdict(int, {"row3": 0, "row4": 0, "row11": 0})
    missing: List[str] = []
    try:
        for ws in wb.worksheets:
            m = SHEET_RE.fullmatch(ws.title.strip())
            if not m:
                continue
            lot = int(m.group(1))
            statuses = []
            for col in range(FIRST_PARTICIPANT_COL, ws.max_column + 1):
                name = ws.cell(row=ROW_NAME, column=col).value
                if not name or not str(name).lower().startswith("ofertant"):
                    continue
                participant = str(name).split(":", 1)[1].strip()
                if not participant:  # «Ofertant:» без имени - лот без участников
                    continue
                stats["participants"] += 1
                rec = offers.get((lot, participant_key(participant)))
                if not rec:
                    missing.append(f"lot {lot}: {participant}")
                    continue
                stats["matched"] += 1
                stats["row3"] += _write(ws, ROW_PRODUCER, col, producer_cell(rec))
                stats["row4"] += _write(ws, ROW_DM, col, dm_cell(rec))
                spec = "\n".join(rec["spec_offered"])
                check = None
                if evidence and spec:
                    check = check_offer(conn, rec["participant"], lot, "; ".join(rec["model"]), spec)
                    stats["evidence"] += bool(check.evidence)
                    stats["no_docs"] += not check.has_docs
                    stats[check.status] += 1
                    spec = f"{spec}\n{check.note()}" if check.note() else spec
                if _write(ws, ROW_SPEC, col, spec):
                    stats["row11"] += 1
                    if check:
                        ws.cell(row=ROW_SPEC, column=col).fill = _fill(check.status)
                        statuses.append(check.status)
            tab = ws.sheet_properties.tabColor
            if statuses and not (tab is not None and str(tab.rgb).upper() in KEEP_TAB_COLORS):
                ws.sheet_properties.tabColor = COLORS[next(s for s in TAB_PRIORITY if s in statuses)]
        wb.save(out_path or result_path)
    finally:
        conn.close()
    stats["missing"] = len(missing)
    if missing:
        logger.info(f"[FILL] Без строки Anexa 22: {', '.join(missing[:20])}{' ...' if len(missing) > 20 else ''}")
    return dict(stats)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Запись данных Anexa 22 в листы лотов result-файла")
    parser.add_argument("result", type=Path)
    parser.add_argument("tender_dir", type=Path)
    parser.add_argument("--out", type=Path, help="куда сохранить (по умолчанию - в тот же файл)")
    parser.add_argument("--no-evidence", action="store_true", help="без поиска подтверждений в даташитах")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="[%(asctime)s] [%(levelname)s] %(message)s")
    sys.stdout.reconfigure(encoding="utf-8")
    stats = fill_result(args.result, args.tender_dir, args.out, evidence=not args.no_evidence)
    print(f"Участников в листах: {stats.get('participants', 0)}, со строкой Anexa 22: {stats.get('matched', 0)}")
    print(f"Записано: строка 3 - {stats.get('row3', 0)}, строка 4 - {stats.get('row4', 0)}, "
          f"строка 11 - {stats.get('row11', 0)} (из них с цитатами из документов: {stats.get('evidence', 0)})")
    print(f"Уровень 2: зелёный - {stats.get(GREEN, 0)}, светло-зелёный - {stats.get(LIGHT_GREEN, 0)}, "
          f"жёлтый - {stats.get(YELLOW, 0)} (из них без технических документов: {stats.get('no_docs', 0)})")
    print(f"Сохранено: {args.out or args.result}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
