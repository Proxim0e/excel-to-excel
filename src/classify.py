"""
Классификация уникальных документов тендера (этап 1, docs/ROADMAP.md).

Каждый уникальный по содержимому файл (documents.sha256 в манифесте) читается один раз:
текст всех страниц сохраняется в document_pages, тип определяется правилами по ключевым словам.

Порядок проверки:
  1. заголовок первой страницы (первые TITLE_CHARS символов) - самый надёжный сигнал;
  2. текст первых двух страниц;
  3. имя файла;
иначе - "other". Правила внутри каждого прохода идут от более специфичных к общим.

Проверка человеком («спросить один раз»): --export выгружает classification_review.xlsx,
в колонке «Правильный тип» исправляются ошибки, --apply переносит их в манифест
(type_source = 'human'); такие решения правила больше не перезаписывают.

Запуск:  python -m src.classify downloads/21641880 [--force] [--export] [--apply [review.xlsx]]
"""
import argparse
import logging
import re
import sqlite3
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Tuple

from src.manifest import MANIFEST_NAME, migrate

logger = logging.getLogger(__name__)

TITLE_CHARS = 600   # «шапка» первой страницы
HEAD_PAGES = 2      # сколько страниц смотрим для классификации

# Тип -> (описание для человека, шаблоны по нормализованному тексту)
# Нормализация: нижний регистр, без диакритики, пробелы схлопнуты.
DOC_TYPES = {
    "anexa22": ("Anexa 22 - specificatii tehnice", [
        r"specificatii? tehnice", r"specificatia tehnica", r"anexa (nr\.? ?)?22\b",
        r"modelul a\w{0,3}icolului", r"specif\w*\.? tehnic\w* deplin", r"tara de origine.{0,40}producator"]),
    "anexa23": ("Anexa 23 - specificatii de pret", [
        r"specificatii? de pret", r"specificatia de pret", r"anexa (nr\.? ?)?23\b",
        r"pret unitar \(fara tva\)", r"suma \(fara tva\)"]),
    "duae": ("DUAE - documentul unic de achizitii european", [
        r"documentul unic de achizitii", r"\bduae\b", r"european single procurement"]),
    "cerere_participare": ("Cerere de participare (anexa 7)", [
        r"cerere(a)? de participare"]),
    "declaratie_valabilitate": ("Declaratie privind valabilitatea ofertei (anexa 8)", [
        r"valabilitatea ofertei"]),
    "garantie_oferta": ("Garantia pentru oferta / scrisoare bancara / ordin de plata", [
        r"garantia? (pentru|de) oferta", r"garantie bancara", r"scrisoare de garantie",
        r"ordin de plata", r"payment order", r"garantia bancara"]),
    "certificat_restante": ("Certificat privind lipsa restantelor (SFS)", [
        r"lipsa (sau (prezenta|existenta) )?restantelor", r"certificat(ul)? fiscal"]),
    "extras_registru": ("Extras din Registrul de stat al persoanelor juridice", [
        r"extras din registrul de stat", r"registrul de stat al persoanelor juridice"]),
    "certificat_inregistrare_firma": ("Certificat de inregistrare a intreprinderii", [
        r"certificat(ul)? de inregistrare a (intreprinderii|intreprinderi|persoanei juridice)",
        r"certificat(ul)? de inregistrare de stat"]),
    "cont_bancar": ("Certificat de atribuire a contului bancar", [
        r"atribuire(a)? (a )?contului", r"contul(ui)? bancar", r"conturi(lor)? bancare",
        r"existenta conturilor", r"conturi(lor)? curente", r"cont curent", r"date bancare"]),
    "mediu_eee": ("Confirmare inscriere in lista producatorilor EEE (Agentia de Mediu)", [
        r"lista producatorilor", r"agentia de mediu", r"\beee\b"]),
    "lista_fondatori": ("Lista fondatorilor / beneficiari efectivi", [
        r"lista fondatorilor", r"fondatorii"]),
    "raport_financiar": ("Situatii financiare / raport financiar", [
        r"situatii(le)? financiare", r"raport(ul)? financiar", r"\bbilant(ul)?\b"]),
    "iso": ("Certificat ISO (13485, 9001, 14001)", [
        r"iso 13485", r"iso 9001", r"iso 14001", r"en iso 1348\d"]),
    "ce_doc": ("Declaratie de conformitate / certificat CE", [
        r"declaration of conformity", r"declaratie de conformitate", r"declaratia de conformitate",
        r"eu declaration", r"ec certificate", r"ce certificate", r"certificat ce\b",
        r"regulation \(eu\) 2017/745", r"93/42/eec"]),
    "amdm": ("Inregistrare AMDM a dispozitivului medical", [
        r"agentia medicamentului", r"registrul de stat al dispozitivelor medicale",
        r"certificat de inregistrare a dispozitivului", r"\bdm0\d{6,}"]),
    "autorizatie_producator": ("Autorizatie / scrisoare a producatorului", [
        r"letter of authori[sz]ation", r"authori[sz]ation letter", r"manufacturer'?s? authori[sz]ation",
        r"scrisoare de autorizare", r"autorizatie(i)? (de la|a) producator", r"authorized distributor"]),
    "procura": ("Procura / imputernicire", [
        r"\bprocura\b", r"imputernicire", r"power of attorney"]),
    "declaratie": ("Alte declaratii ale ofertantului", [
        r"\bdeclaratie\b", r"\bdeclaratia\b", r"\bdeclar pe propria raspundere", r"declaratii"]),
    "datasheet": ("Fisa tehnica / catalog / brosura / instructiuni", [
        r"data ?sheet", r"technical data", r"\bcatalog", r"brochure", r"brosura", r"instructions for use",
        r"user manual", r"instructiuni de utilizare", r"\bprospect", r"fisa tehnica", r"specifications?\b",
        r"product information"]),
}
OTHER = "other"
NO_TEXT = "no_text"

# Только по имени файла (в тексте такие шаблоны дают ложные срабатывания).
# Имя нормализуется: знаки препинания и цифры -> пробелы.
FILENAME_ONLY = {
    "amdm": [r"\blot(ul)?\b.*\breg\b", r"\bdm\b.*\binregistr"],
    "datasheet": [r"brosur", r"brochure", r"catalog", r"prospect", r"\bimagin", r"datasheet"],
}

# Штамп электронной подписи MoldSign: в сканах это единственный текст на странице,
# а в остальных документах он занимает «шапку». Вырезается перед классификацией.
SIGNATURE_RE = re.compile(r"digitally signed by .{0,160}?location: ?[a-z]+( europeana| moldova)*")
MIN_TEXT_CHARS = 40  # меньше - считаем, что текстового слоя нет (скан, картинка)

# Признаки страницы каталога / даташита, когда явных слов «datasheet» нет
PRODUCT_MARKERS = [r"www\.", r"[®©™]", r"\bref\.?\b", r"\bp/n\b", r"\bcode\b", r"\bmodel\b",
                   r"\d+ ?mm\b", r"\d+ ?ml\b", r"\bdisposable\b", r"\breusable\b", r"\bsterile\b"]
_PRODUCT_RX = [re.compile(p) for p in PRODUCT_MARKERS]

# Вне «шапки» общие слова вроде «specificatii tehnice» встречаются в декларациях,
# поэтому в тексте Anexa 22/23 узнаются только по заголовкам колонок таблицы.
STRUCTURAL = {
    "anexa22": [r"modelul a\w{0,3}icolului", r"specif\w*\.? tehnic\w* deplin", r"specificatia tehnica propusa",
                r"tara de origine.{0,40}producator"],
    "anexa23": [r"pret unitar \(fara tva\)", r"suma \(fara tva\)"],
}
# Таблицы Anexa продолжаются на следующих страницах без заголовка - страница наследует тип
CONTINUED_TYPES = ("anexa22", "anexa23")

# Порядок типов в DOC_TYPES важен: Anexa 22/23 и DUAE проверяются раньше «declaratie»,
# потому что в их шапках встречается слово «declaratie».
ORDER = list(DOC_TYPES)
_COMPILED = {t: [re.compile(p) for p in pats] for t, (_, pats) in DOC_TYPES.items()}


def _relaxed(pattern: str) -> str:
    """Шаблон для имён файлов без пробелов: «certificatprivindlipsarestantelor»."""
    return pattern.replace(r"\b", "").replace(" ", " ?")


_FILENAME_COMPILED = {t: [re.compile(_relaxed(p)) for p in pats] + [re.compile(p) for p in FILENAME_ONLY.get(t, [])]
                      for t, (_, pats) in DOC_TYPES.items()}
_TEXT_COMPILED = {t: ([re.compile(p) for p in STRUCTURAL[t]] if t in STRUCTURAL else rxs)
                  for t, rxs in _COMPILED.items()}


@dataclass
class Classification:
    doc_type: str
    confidence: float
    rule: str
    page_types: List[Optional[str]] = field(default_factory=list)  # None - на странице нет текста
    has_text: bool = True

    @property
    def contains(self) -> List[str]:
        """Типы разделов в порядке появления (для PDF-«пачек»)."""
        seen = []
        for t in self.page_types:
            if t and t not in seen:
                seen.append(t)
        return seen


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKD", text or "")
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text.lower()).strip()


def clean_page(text: str) -> str:
    """Нормализованный текст страницы без штампов подписи."""
    return SIGNATURE_RE.sub(" ", normalize(text)).strip()


def normalize_filename(name: str) -> str:
    name = normalize(name)
    name = re.sub(r"\.(pdf|signed|semnat)\b", " ", name)
    return re.sub(r"[\W\d_]+", " ", name).strip()


def _first_match(text: str, compiled=None) -> Optional[Tuple[str, str]]:
    compiled = compiled or _COMPILED
    for doc_type in ORDER:
        for rx in compiled[doc_type]:
            if rx.search(text):
                return doc_type, rx.pattern
    return None


def classify_pages(cleaned: List[str]) -> List[Optional[str]]:
    """Тип каждой страницы по её «шапке»; продолжение таблицы Anexa наследует тип."""
    result, prev = [], None
    for text in cleaned:
        if len(text) < MIN_TEXT_CHARS:
            result.append(None)
            continue
        hit = _first_match(text[:TITLE_CHARS])
        page_type = hit[0] if hit else (prev if prev in CONTINUED_TYPES else OTHER)
        result.append(page_type)
        prev = page_type
    return result


def classify(pages: List[str], filename: str = "") -> Classification:
    """Тип документа (по первой странице с текстом) и типы всех страниц."""
    cleaned = [clean_page(p) for p in pages]
    texted = [p for p in cleaned if len(p) >= MIN_TEXT_CHARS]
    fname = normalize_filename(filename)

    if not texted:
        hit = _first_match(fname, _FILENAME_COMPILED)
        page_types = [None] * len(pages)
        if hit:
            return Classification(hit[0], 0.3, f"filename(no text):{hit[1]}", page_types, False)
        return Classification(NO_TEXT, 0.0, "no text layer", page_types, False)

    page_types = classify_pages(cleaned)
    title = texted[0][:TITLE_CHARS]
    head = " ".join(texted[:HEAD_PAGES])
    for region, compiled, conf, label in ((title, _COMPILED, 0.9, "title"), (head, _TEXT_COMPILED, 0.6, "text"),
                                          (fname, _FILENAME_COMPILED, 0.4, "filename")):
        hit = _first_match(region, compiled)
        if hit:
            doc_type, rule = hit
            # Один PDF с обеими анексами: техническая и ценовая части вместе
            if doc_type in ("anexa22", "anexa23") and {"anexa22", "anexa23"} <= set(page_types):
                return Classification("anexa22+23", conf, f"{label}:{rule} + pages", page_types)
            return Classification(doc_type, conf, f"{label}:{rule}", page_types)

    markers = sum(1 for rx in _PRODUCT_RX if rx.search(head))
    if markers >= 2:
        return Classification("datasheet", 0.4, f"product markers x{markers}", page_types)
    return Classification(OTHER, 0.0, "no rule", page_types)


def extract_pages(path: Path) -> List[str]:
    import fitz  # PyMuPDF: импорт здесь, чтобы манифест и тесты правил работали без него
    with fitz.open(path) as doc:
        return [page.get_text() for page in doc]


def classify_tender(tender_dir: Path, force: bool = False) -> dict:
    """Классифицирует уникальные документы манифеста, у которых ещё нет типа (или все при force)."""
    db = Path(tender_dir) / MANIFEST_NAME
    conn = sqlite3.connect(db)
    migrate(conn)
    where = "" if force else "WHERE d.doc_type IS NULL"
    rows = conn.execute(
        "SELECT d.sha256, MIN(r.path), MIN(r.original_name) FROM documents d "
        f"JOIN document_refs r USING (sha256) {where} GROUP BY d.sha256").fetchall()

    counts = {}
    now = datetime.now().isoformat(timespec="seconds")
    for sha, path, original_name in rows:
        # Решение человека правилами не перезаписываем
        if conn.execute("SELECT type_source FROM documents WHERE sha256 = ?", (sha,)).fetchone()[0] == "human":
            continue
        try:
            pages = extract_pages(Path(path))
        except Exception as e:
            logger.warning(f"[CLASSIFY] Не удалось прочитать {path}: {type(e).__name__} - {e}")
            pages = []
        c = classify(pages, original_name or Path(path).name)
        conn.execute("DELETE FROM document_pages WHERE sha256 = ?", (sha,))
        conn.executemany("INSERT INTO document_pages (sha256, page_no, text, page_type) VALUES (?, ?, ?, ?)",
                         [(sha, i + 1, t, pt) for i, (t, pt) in enumerate(zip(pages, c.page_types))])
        conn.execute(
            "UPDATE documents SET doc_type = ?, type_confidence = ?, type_source = 'rule', type_rule = ?, "
            "pages = ?, has_text = ?, contains = ?, classified_at = ? WHERE sha256 = ?",
            (c.doc_type, c.confidence, c.rule, len(pages), int(c.has_text), ",".join(c.contains), now, sha))
        counts[c.doc_type] = counts.get(c.doc_type, 0) + 1
    conn.commit()
    conn.close()
    return counts


def type_summary(tender_dir: Path) -> List[Tuple[str, int, int]]:
    """(тип, уникальных документов, мест в лотах) по всему манифесту."""
    conn = sqlite3.connect(Path(tender_dir) / MANIFEST_NAME)
    try:
        return conn.execute(
            "SELECT COALESCE(d.doc_type, '?'), COUNT(DISTINCT d.sha256), COUNT(r.path) "
            "FROM documents d JOIN document_refs r USING (sha256) "
            "GROUP BY 1 ORDER BY 2 DESC").fetchall()
    finally:
        conn.close()


REVIEW_NAME = "classification_review.xlsx"
REVIEW_SHEET = "Документы"
REVIEW_COLUMN = "Правильный тип"
EXTRA_TYPES = {
    "anexa22+23": "Anexa 22 si 23 intr-un singur fisier",
    OTHER: "Altele - nu se proceseaza",
    NO_TEXT: "Fara text (scanare) - tipul nu poate fi determinat",
}


def all_types() -> dict:
    types = {t: desc for t, (desc, _) in DOC_TYPES.items()}
    types.update(EXTRA_TYPES)
    return types


def export_review(tender_dir: Path, xlsx_path: Optional[Path] = None) -> Path:
    """Таблица для проверки человеком: один уникальный документ - одна строка, сначала неуверенные."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    def safe(value):
        return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value

    xlsx_path = Path(xlsx_path or Path(tender_dir) / REVIEW_NAME)
    conn = sqlite3.connect(Path(tender_dir) / MANIFEST_NAME)
    rows = conn.execute(
        "SELECT d.sha256, d.doc_type, d.type_confidence, d.type_source, d.type_rule, d.contains, d.pages, "
        "d.has_text, GROUP_CONCAT(DISTINCT r.participant), COUNT(DISTINCT r.lot_no), MIN(r.original_name), "
        "(SELECT text FROM document_pages p WHERE p.sha256 = d.sha256 AND p.page_no = 1), MIN(r.path) "
        "FROM documents d JOIN document_refs r USING (sha256) GROUP BY d.sha256 "
        "ORDER BY d.type_source = 'human', d.type_confidence, d.doc_type").fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = REVIEW_SHEET
    header = ["sha256", "Тип (система)", "Уверенность", "Источник", "Правило", "Внутри файла", "Страниц",
              "Есть текст", "Участники", "Лотов", "Имя файла", "Начало первой страницы", "Путь", REVIEW_COLUMN]
    ws.append(header)
    for (sha, doc_type, conf, source, rule, contains, pages, has_text, parts, lots, name,
         first_page, path) in rows:
        ws.append([safe(v) for v in (sha, doc_type, conf, source, rule, contains, pages,
                                     "да" if has_text else "нет", parts, lots, name,
                                     clean_page(first_page or "")[:300], path,
                                     doc_type if source == "human" else None)])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws[f"{chr(64 + len(header))}1"].fill = PatternFill("solid", fgColor="FFFF00")
    for col, width in zip("ABCDEFGHIJKLMN", (14, 22, 11, 9, 30, 30, 8, 9, 30, 7, 40, 80, 20, 22)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    types_ws = wb.create_sheet("Типы")
    types_ws.append(["Тип", "Описание"])
    for t, desc in all_types().items():
        types_ws.append([t, desc])
    types_ws.column_dimensions["A"].width = 30
    types_ws.column_dimensions["B"].width = 70

    dv = DataValidation(type="list", formula1=f"=Типы!$A$2:$A${len(all_types()) + 1}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"N2:N{len(rows) + 1}")

    wb.save(xlsx_path)
    return xlsx_path


def apply_review(tender_dir: Path, xlsx_path: Optional[Path] = None) -> Tuple[int, List[str]]:
    """Переносит исправления человека в манифест. Возвращает (сколько применено, ошибки)."""
    from openpyxl import load_workbook

    xlsx_path = Path(xlsx_path or Path(tender_dir) / REVIEW_NAME)
    wb = load_workbook(xlsx_path, read_only=True)
    try:
        # read_only держит файл открытым до close(): читаем всё сразу и отпускаем файл
        rows = iter(list(wb[REVIEW_SHEET].iter_rows(values_only=True)))
    finally:
        wb.close()
    header = list(next(rows))
    sha_i, fix_i = header.index("sha256"), header.index(REVIEW_COLUMN)
    valid = all_types()

    conn = sqlite3.connect(Path(tender_dir) / MANIFEST_NAME)
    applied, errors = 0, []
    now = datetime.now().isoformat(timespec="seconds")
    for n, row in enumerate(rows, start=2):
        sha, fix = row[sha_i], (row[fix_i] or "").strip() if row[fix_i] else ""
        if not fix:
            continue
        if fix not in valid:
            errors.append(f"строка {n}: неизвестный тип {fix!r}")
            continue
        cur = conn.execute(
            "UPDATE documents SET doc_type = ?, type_confidence = 1.0, type_source = 'human', "
            "type_rule = 'human', classified_at = ? WHERE sha256 = ?", (fix, now, sha))
        if cur.rowcount:
            applied += 1
        else:
            errors.append(f"строка {n}: документ {str(sha)[:12]} не найден в манифесте")
    conn.commit()
    conn.close()
    return applied, errors


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Классификация документов тендера по манифесту")
    parser.add_argument("tender_dir", type=Path)
    parser.add_argument("--force", action="store_true", help="переклассифицировать всё, кроме решений человека")
    parser.add_argument("--export", action="store_true", help=f"выгрузить {REVIEW_NAME} для проверки")
    parser.add_argument("--apply", nargs="?", const="", metavar="XLSX",
                        help=f"применить колонку «{REVIEW_COLUMN}» (по умолчанию {REVIEW_NAME} в папке тендера)")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="[%(asctime)s] [%(levelname)s] %(message)s")
    sys.stdout.reconfigure(encoding="utf-8")

    if args.apply is not None:
        applied, errors = apply_review(args.tender_dir, Path(args.apply) if args.apply else None)
        print(f"Исправлений применено: {applied}")
        for e in errors:
            print(f"  ! {e}")

    counts = classify_tender(args.tender_dir, force=args.force)
    print(f"Классифицировано сейчас: {sum(counts.values())}")
    for doc_type, docs, refs in type_summary(args.tender_dir):
        print(f"  {doc_type:30} {docs:5} уникальных  {refs:6} в лотах")

    if args.export:
        try:
            print(f"Таблица для проверки: {export_review(args.tender_dir)}")
        except PermissionError:
            print(f"Не удалось записать {REVIEW_NAME}: файл открыт в Excel? Закройте его и повторите.")
            return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
