"""
Эталон для оценки разбора Anexa 22: то, что человек вручную перенёс в листы лотов.

Раскладка листа лота (участник - колонка начиная с E):
  строка 2  - «Ofertant: <участник>»
  строка 3  - «Producator: X / Țara: Y / Model: Z» (через переводы строк)
  строка 4  - номер DM и прочие регистрации
  строка 11 - предложенная спецификация
Заполнено не у всех участников, поэтому эталон - только непустые ячейки.
Колонки «clarificări» (уточнения) вставляются справа от участника без «Ofertant:» и пропускаются.
"""
import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

FIRST_PARTICIPANT_COL = 5  # E
SHEET_RE = re.compile(r"lot\s*(\d+)", re.IGNORECASE)
DM_RE = re.compile(r"\bDM\s?0\d{6,}", re.IGNORECASE)


@dataclass
class GoldOffer:
    lot: int
    sheet: str
    participant: str
    producer: Optional[str]
    country: Optional[str]
    model: Optional[str]
    dm: List[str]
    dm_raw: Optional[str]
    spec: Optional[str]


def _field(text: str, label: str, labels: str) -> Optional[str]:
    """Значение после «Label:» до следующей метки (переводы строк внутри значения допустимы)."""
    m = re.search(rf"{label}\s*:\s*(.*?)(?=(?:{labels})\s*:|$)", text, re.IGNORECASE | re.DOTALL)
    if not m:
        return None
    value = re.sub(r"\s+", " ", m.group(1)).strip()
    return value or None


def parse_row3(text: str):
    labels = r"Producator|Producătorul|Producător|Țara|Ţara|Tara|Model"
    text = str(text or "")
    return (_field(text, r"Produc[aă]tor(?:ul)?", labels),
            _field(text, r"[ȚŢT]ara", labels),
            _field(text, r"Model", labels))


def dm_numbers(text) -> List[str]:
    return [re.sub(r"\s", "", m).upper() for m in DM_RE.findall(str(text or ""))]


def read_gold(xlsx_path: Path) -> List[GoldOffer]:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        offers = []
        for name in wb.sheetnames:
            m = SHEET_RE.fullmatch(name.strip()) or SHEET_RE.match(name.strip())
            if not m:
                continue
            ws = wb[name]
            # В read_only пустые ячейки не знают свою строку - нумеруем строки сами
            rows = dict(enumerate(ws.iter_rows(min_row=1, max_row=11), start=1))
            if 2 not in rows:
                continue
            for col, cell in enumerate(rows[2], start=1):
                value = getattr(cell, "value", None)
                if col < FIRST_PARTICIPANT_COL or not value or not str(value).strip().lower().startswith("ofertant"):
                    continue
                participant = str(value).split(":", 1)[1].strip()
                if not participant:
                    continue

                def at(row):
                    r = rows.get(row)
                    return getattr(r[col - 1], "value", None) if r and len(r) >= col else None

                row3, row4, row11 = at(3), at(4), at(11)
                if not any((row3, row4, row11)):
                    continue
                producer, country, model = parse_row3(row3)
                if not any((producer, country, model, row4, row11)):
                    continue
                offers.append(GoldOffer(
                    lot=int(m.group(1)), sheet=name, participant=participant,
                    producer=producer, country=country, model=model,
                    dm=dm_numbers(row4), dm_raw=str(row4).strip() if row4 else None,
                    spec=re.sub(r"\s+", " ", str(row11)).strip() if row11 else None))
        return offers
    finally:
        wb.close()
