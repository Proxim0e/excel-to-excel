"""
python -m src.anexa22 downloads/21641880 [--gold "21641880 CDM27 B.xlsx"] [--errors N]

Разбирает страницы Anexa 22 из манифеста тендера, сохраняет строки в anexa22_offers
и, если задан эталон, печатает точность по полям.
"""
import argparse
import logging
import sys
from pathlib import Path

from src.anexa22.evaluate import FIELDS, compare
from src.anexa22.gold import read_gold
from src.anexa22.parse import parse_tender, save_offers


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Разбор Anexa 22 тендера")
    parser.add_argument("tender_dir", type=Path)
    parser.add_argument("--gold", type=Path, help="Excel тендера, заполненный человеком (эталон)")
    parser.add_argument("--errors", type=int, default=30, help="сколько расхождений показать")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.WARNING)
    sys.stdout.reconfigure(encoding="utf-8")

    offers = parse_tender(args.tender_dir)
    save_offers(args.tender_dir, offers)
    with_lot = [o for o in offers if o.lot is not None]
    print(f"Строк Anexa 22: {len(offers)}, с номером лота: {len(with_lot)}, "
          f"участников: {len({o.participant for o in offers})} -> таблица anexa22_offers манифеста")
    if not args.gold:
        return 0

    r = compare(read_gold(args.gold), offers)
    print(f"\nЭталон: {r['gold']} оферт, найдено строк Anexa 22: {r['found']} ({r['found'] / r['gold']:.1%})")
    for name in FIELDS:
        ok, n = r["fields"][name]
        print(f"  {name:10} {ok:4}/{n:<4} {ok / n:.1%}" if n else f"  {name:10} -")
    print("\nПо участникам (найдено строк / верных полей):")
    for who, p in sorted(r["per_participant"].items(), key=lambda kv: kv[1]["ok"] / max(kv[1]["checked"], 1)):
        share = p["ok"] / p["checked"] if p["checked"] else 0
        print(f"  {who[:28]:28} {p['found']:3}/{p['gold']:<3}  поля {p['ok']:3}/{p['checked']:<3} {share:.0%}")
    if r["misses"]:
        print("\nНе найдены (участник, лот):", [(m.participant[:18], m.lot) for m in r["misses"][:args.errors]])
    print("\nРасхождения (участник, лот, поле, эталон | распознано):")
    for who, lot, name, expected, got, _ in r["errors"][:args.errors]:
        print(f"  {who[:18]:18} {lot:3} {name:8} {str(expected)[:60]!r:62} | {str(got)[:60]!r}")
    report = write_report(args.tender_dir, r["errors"])
    print(f"\nОтчёт расхождений: {report}")
    return 0


REPORT_NAME = "anexa22_vs_excel.xlsx"


def write_report(tender_dir: Path, errors) -> Path:
    """Расхождения с эталоном для проверки человеком: где ошибка - в разборе или в Excel."""
    import sqlite3

    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font

    from src.manifest import MANIFEST_NAME

    conn = sqlite3.connect(Path(tender_dir) / MANIFEST_NAME)
    paths = dict(conn.execute("SELECT sha256, MIN(path) FROM document_refs GROUP BY sha256"))
    conn.close()

    def safe(v):
        v = ", ".join(v) if isinstance(v, (list, tuple, set)) else v
        return ILLEGAL_CHARACTERS_RE.sub("", str(v)) if v not in (None, "") else ""

    wb = Workbook()
    ws = wb.active
    ws.title = "Расхождения"
    ws.append(["Участник", "Лот", "Поле", "В Excel (эталон)", "В Anexa 22 (распознано)", "Страница", "Файл",
               "Кто прав? (excel / pdf / оба неверно)"])
    for who, lot, name, expected, got, offer in sorted(errors, key=lambda e: (e[2], e[0], e[1])):
        ws.append([safe(who), lot, name, safe(expected), safe(got), offer.page_no,
                   safe(paths.get(offer.sha256)), ""])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFGH", (24, 6, 10, 60, 60, 9, 50, 22)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row[3:5]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    path = Path(tender_dir) / REPORT_NAME
    wb.save(path)
    return path


if __name__ == "__main__":
    sys.exit(main())
