"""
Отчёт проверки AMDM в Excel.

Правило из концепции: система не трогает твои ячейки. Результат пишется на отдельный лист
«Verificare AMDM» (одна строка на участника и номер DM) в КОПИЮ книги; исходный файл не меняется.
"""
from __future__ import annotations

from typing import List

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.amdm.models import CheckResult, OK, CHECK, FAIL, NO_DM, IN_PROGRESS

SHEET = "Verificare AMDM"
FILLS = {
    OK: PatternFill("solid", start_color="FF00B050", end_color="FF00B050"),          # зелёный, как в кодбуке
    CHECK: PatternFill("solid", start_color="FFFFFF00", end_color="FFFFFF00"),       # жёлтый
    FAIL: PatternFill("solid", start_color="FFFF0000", end_color="FFFF0000"),        # красный
    NO_DM: PatternFill("solid", start_color="FFFFC000", end_color="FFFFC000"),       # оранжевый
    IN_PROGRESS: PatternFill("solid", start_color="FFD9D9D9", end_color="FFD9D9D9"), # серый
}
HEADERS = ["Lot (foaie)", "Ofertant", "Producător (ofertă)", "Model (ofertă)", "Nr. DM",
           "Status", "Observații",
           "Producător (registru)", "Model (registru)", "Nr. catalog (registru)", "Denumire (registru)",
           "Ordin / Data (registru)", "Rândul 4 (original)"]
WIDTHS = [12, 26, 28, 28, 15, 12, 60, 28, 24, 24, 34, 22, 30]


def write_report(wb, results: List[CheckResult]) -> None:
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET, 0)
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    priority = {FAIL: 0, CHECK: 1, NO_DM: 2, IN_PROGRESS: 3, OK: 4}
    for r in sorted(results, key=lambda r: (priority[r.status], _lot_key(r.offer.sheet), r.offer.column)):
        rec = r.record
        ws.append([
            r.offer.sheet, r.offer.supplier, r.offer.producer, r.offer.model, r.dm,
            r.status, "\n".join(r.reasons),
            rec.producer if rec else None, rec.model if rec else None,
            rec.catalog if rec else None, rec.name if rec else None,
            f"{rec.order or ''} {rec.order_date or ''}".strip() if rec else None,
            r.offer.raw_row4.strip() or None,
        ])
        row = ws.max_row
        for c in ws[row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=6).fill = FILLS[r.status]
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"


def _lot_key(sheet: str):
    import re
    m = re.search(r'\d+', sheet)
    return (int(m.group(0)) if m else 10**9, sheet)
