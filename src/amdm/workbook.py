"""
Чтение заявленных регистраций из Excel оценки (листы 'lot N').

Колонка участника = ячейка строки 2 вида 'Ofertant: <имя>'. Колонки уточнений (clarificări),
вставленные справа от участника, в строке 2 'Ofertant' не содержат и сюда не попадают.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import List

from src.amdm.models import OfferRegistration
from src.amdm.parse import parse_row3, parse_row4

LOT_SHEET_RE = re.compile(r'^\s*lot\s*\d', re.I)
SUPPLIER_RE = re.compile(r'^\s*ofertant\s*:\s*', re.I)


def load_workbook_tolerant(path, **kwargs):
    """
    Открывает книгу, даже если из неё удалили картинки (папку xl/media) ради размера:
    обычный openpyxl в этом случае падает с KeyError на ссылке на изображение.
    """
    import openpyxl.reader.excel as excel_reader
    from openpyxl import load_workbook

    original = excel_reader.find_images

    def safe_find_images(archive, path_):
        try:
            return original(archive, path_)
        except KeyError:
            return [], []

    excel_reader.find_images = safe_find_images
    try:
        return load_workbook(path, **kwargs)
    finally:
        excel_reader.find_images = original


def read_offers(wb) -> List[OfferRegistration]:
    offers: List[OfferRegistration] = []
    for ws in wb.worksheets:
        if not LOT_SHEET_RE.match(ws.title) or ws.title.strip().lower() == "lot x":
            continue
        for cell in ws[2]:
            value = cell.value
            if not isinstance(value, str) or not SUPPLIER_RE.match(value):
                continue
            supplier = SUPPLIER_RE.sub("", value).strip()
            if not supplier:
                continue                      # пустая колонка шаблона
            row3 = ws.cell(row=3, column=cell.column).value
            row4 = ws.cell(row=4, column=cell.column).value
            p3, p4 = parse_row3(row3), parse_row4(row4)
            offers.append(OfferRegistration(
                sheet=ws.title, column=cell.column, supplier=supplier,
                producer=p3["producer"], country=p3["country"], model=p3["model"],
                dm_numbers=p4["dm_numbers"], order=p4["order"], order_date=p4["order_date"],
                quality_docs=p4["quality_docs"], note=p4["note"],
                raw_row3=str(row3 or ""), raw_row4=str(row4 or ""),
            ))
    return offers


def read_offers_from_file(path: Path) -> List[OfferRegistration]:
    return read_offers(load_workbook_tolerant(path, data_only=True))
